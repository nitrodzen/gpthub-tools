from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
import time
import uuid
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from datetime import time as datetime_time
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from charset_normalizer import from_bytes
from defusedxml import ElementTree as DefusedET
from defusedxml.common import DefusedXmlException
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.worksheet.hyperlink import Hyperlink
from PIL import Image, UnidentifiedImageError

from .models import ErrorCode, JobFailure, JobWarning

MAX_SPREADSHEET_SHEETS = 50
MAX_SPREADSHEET_NONEMPTY_CELLS = 100_000
MAX_WORD_TABLES = 200
MAX_WORD_TABLE_CELLS = 100_000
MAX_EXCEL_COLUMNS = 16_384
MAX_EXCEL_CELL_CHARACTERS = 32_767
MAX_IMAGES = 50
MAX_MEDIA_BYTES = 100 * 1024 * 1024
MAX_OUTPUT_TABLE_CELLS = 100_000
MAX_WORD_COLUMNS_PER_BAND = 8
MAX_WORD_PARAGRAPHS = 200_000
MAX_OFFICE_XML_TEXT_CHARS = 20_000_000
MAX_OFFICE_XML_ATTRIBUTE_CHARS = 10_000_000
MAX_OFFICE_XML_ELEMENTS = 1_500_000
MAX_OFFICE_XML_DEPTH = 256
MAX_OFFICE_STYLE_RECORDS = 50_000
MAX_OFFICE_RELATIONSHIPS = 20_000
MAX_OFFICE_SHARED_STRINGS = 100_000
MAX_EMBEDDED_IMAGE_PIXELS = 40_000_000
MAX_TOTAL_IMAGE_PIXELS = 100_000_000

SUPPORTED_WORD_INPUTS = {".doc", ".docx", ".odt", ".rtf"}
SUPPORTED_EXCEL_INPUTS = {".xls", ".xlsx", ".ods", ".csv"}
RASTER_CONTENT_TYPES = {
    "image/bmp",
    "image/gif",
    "image/jpeg",
    "image/png",
    "image/tiff",
    "image/webp",
}


@dataclass(slots=True)
class CellData:
    value: str = ""
    bold: bool = False
    italic: bool = False
    font_name: str | None = None
    font_size: float | None = None
    font_color: str | None = None
    fill: str | None = None
    alignment: str | None = None
    border: bool = False
    hyperlink: str | None = None
    comment: tuple[str, str] | None = None


@dataclass(slots=True)
class RasterPlacement:
    data: bytes
    row: int
    column: int
    width: int
    height: int


@dataclass(slots=True)
class MediaBudget:
    images: int = 0
    bytes: int = 0
    pixels: int = 0

    def add(self, blob: bytes, width: int, height: int) -> None:
        self.images += 1
        self.bytes += len(blob)
        self.pixels += width * height
        if self.images > MAX_IMAGES:
            raise _too_complex("images", self.images, MAX_IMAGES)
        if self.bytes > MAX_MEDIA_BYTES:
            raise _too_complex("mediaBytes", self.bytes, MAX_MEDIA_BYTES)
        if self.pixels > MAX_TOTAL_IMAGE_PIXELS:
            raise _too_complex("imagePixels", self.pixels, MAX_TOTAL_IMAGE_PIXELS)


@dataclass(slots=True)
class SheetData:
    name: str
    rows: list[list[CellData]]
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)
    images: list[RasterPlacement] = field(default_factory=list)
    hidden: bool = False
    source_columns: list[int] = field(default_factory=list)
    column_widths: dict[int, float] = field(default_factory=dict)


@dataclass(slots=True)
class ParsedOffice:
    sheets: list[SheetData]
    formulas: bool = False
    unsupported_objects: bool = False


@dataclass(frozen=True, slots=True)
class OperationResult:
    path: Path
    warnings: list[JobWarning] = field(default_factory=list)


WARNING_MESSAGES = {
    "en": {
        "FORMULAS_AS_VALUES": (
            "Formulas were added as saved displayed values, or as formula text "
            "when no saved value existed."
        ),
        "OBJECTS_SKIPPED": "Unsupported content or embedded objects were skipped.",
        "LAYOUT_SIMPLIFIED": "The source layout was simplified to keep the result editable.",
        "CSV_DETECTION_GUESSED": "CSV encoding or delimiter was detected automatically.",
    },
    "ru": {
        "FORMULAS_AS_VALUES": (
            "Формулы перенесены как сохранённые отображаемые значения, "
            "а при их отсутствии — как текст формулы."
        ),
        "OBJECTS_SKIPPED": "Неподдерживаемое содержимое или встроенные объекты пропущены.",
        "LAYOUT_SIMPLIFIED": "Макет источника упрощён, чтобы результат оставался редактируемым.",
        "CSV_DETECTION_GUESSED": "Кодировка или разделитель CSV определены автоматически.",
    },
}

DOCUMENT_LABELS = {
    "en": {
        "columns": "Columns",
        "hidden": "hidden",
        "no_nonempty_sheets": "No non-empty sheets.",
    },
    "ru": {
        "columns": "Столбцы",
        "hidden": "скрыт",
        "no_nonempty_sheets": "Нет непустых листов.",
    },
}


def _locale(options: dict[str, Any]) -> str:
    locale = str(options.get("locale", "en")).lower()
    if locale not in {"ru", "en"}:
        raise JobFailure(ErrorCode.INVALID_FILE, "Locale must be ru or en")
    return locale


def _conversion_deadline(options: dict[str, Any]) -> float | None:
    value = options.get("_deadlineMonotonic")
    return float(value) if isinstance(value, int | float) else None


def _warning(code: str, locale: str, details: dict[str, Any] | None = None) -> JobWarning:
    return JobWarning(code=code, message=WARNING_MESSAGES[locale][code], details=details)


def _deduplicate_warnings(warnings: list[JobWarning]) -> list[JobWarning]:
    result: list[JobWarning] = []
    codes: set[str] = set()
    for warning in warnings:
        if warning.code not in codes:
            codes.add(warning.code)
            result.append(warning)
    return result


def _clean_stem(name: str) -> str:
    cleaned = re.sub(r"[^\w.-]+", "_", Path(name).stem, flags=re.UNICODE).strip("._")
    return cleaned[:80] or "result"


def _safe_hyperlink(value: str | None) -> str | None:
    if not value:
        return None
    parsed = urlparse(value)
    if parsed.scheme.casefold() not in {"http", "https", "mailto"}:
        return None
    return value


def _too_complex(kind: str, value: int, limit: int) -> JobFailure:
    return JobFailure(
        ErrorCode.OFFICE_TOO_COMPLEX,
        "The Office file is too complex to convert safely",
        {"kind": kind, "value": value, "limit": limit},
    )


def _check_deadline(deadline: float | None) -> None:
    if deadline is not None and time.monotonic() >= deadline:
        raise JobFailure(ErrorCode.TIMEOUT, "Office conversion timed out")


def _remaining_timeout(deadline: float | None, maximum: float = 600) -> float:
    if deadline is None:
        return maximum
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise JobFailure(ErrorCode.TIMEOUT, "Office conversion timed out")
    return max(0.1, min(maximum, remaining))


def _safe_xml_package(path: Path, deadline: float | None = None) -> None:
    """Reject XML entities/DTDs before python-docx/openpyxl parse the package."""

    try:
        word_tables = 0
        word_cells = 0
        word_paragraphs = 0
        word_table_grids: list[int] = []
        word_table_rows: list[int | None] = []
        spreadsheet_sheets = 0
        spreadsheet_cells = 0
        spreadsheet_materialized_cells = 0
        spreadsheet_merged_cells = 0
        text_characters = 0
        attribute_characters = 0
        parsed_elements = 0
        style_records = 0
        relationships = 0
        shared_strings = 0
        with zipfile.ZipFile(path) as archive:
            for info in archive.infolist():
                _check_deadline(deadline)
                if info.filename.casefold().endswith((".xml", ".rels")):
                    with archive.open(info) as source:
                        depth = 0
                        for event, element in DefusedET.iterparse(source, events=("start", "end")):
                            local_name = element.tag.rsplit("}", 1)[-1]
                            lowered_name = info.filename.casefold()
                            if event == "start":
                                depth += 1
                                parsed_elements += 1
                                attribute_characters += sum(
                                    len(str(name)) + len(str(value))
                                    for name, value in element.attrib.items()
                                )
                                if parsed_elements > MAX_OFFICE_XML_ELEMENTS:
                                    raise _too_complex(
                                        "officeXmlElements",
                                        parsed_elements,
                                        MAX_OFFICE_XML_ELEMENTS,
                                    )
                                if depth > MAX_OFFICE_XML_DEPTH:
                                    raise _too_complex(
                                        "officeXmlDepth",
                                        depth,
                                        MAX_OFFICE_XML_DEPTH,
                                    )
                                if attribute_characters > MAX_OFFICE_XML_ATTRIBUTE_CHARS:
                                    raise _too_complex(
                                        "officeXmlAttributeCharacters",
                                        attribute_characters,
                                        MAX_OFFICE_XML_ATTRIBUTE_CHARS,
                                    )
                                if lowered_name == "xl/styles.xml" and local_name in {
                                    "border",
                                    "dxf",
                                    "fill",
                                    "font",
                                    "xf",
                                }:
                                    style_records += 1
                                    if style_records > MAX_OFFICE_STYLE_RECORDS:
                                        raise _too_complex(
                                            "officeStyleRecords",
                                            style_records,
                                            MAX_OFFICE_STYLE_RECORDS,
                                        )
                                elif lowered_name == "xl/sharedstrings.xml" and local_name == "si":
                                    shared_strings += 1
                                    if shared_strings > MAX_OFFICE_SHARED_STRINGS:
                                        raise _too_complex(
                                            "officeSharedStrings",
                                            shared_strings,
                                            MAX_OFFICE_SHARED_STRINGS,
                                        )
                                elif (
                                    lowered_name.endswith(".rels") and local_name == "Relationship"
                                ):
                                    relationships += 1
                                    if relationships > MAX_OFFICE_RELATIONSHIPS:
                                        raise _too_complex(
                                            "officeRelationships",
                                            relationships,
                                            MAX_OFFICE_RELATIONSHIPS,
                                        )
                                if parsed_elements % 1024 == 0:
                                    _check_deadline(deadline)
                                if lowered_name == "word/document.xml":
                                    if local_name == "tbl":
                                        word_table_grids.append(0)
                                        word_table_rows.append(None)
                                    elif local_name == "gridCol" and word_table_grids:
                                        word_table_grids[-1] += 1
                                        if word_table_grids[-1] > MAX_EXCEL_COLUMNS:
                                            raise _too_complex(
                                                "wordTableColumns",
                                                word_table_grids[-1],
                                                MAX_EXCEL_COLUMNS,
                                            )
                                    elif local_name == "tr" and word_table_rows:
                                        word_table_rows[-1] = 0
                                    elif (
                                        local_name == "tc"
                                        and word_table_rows
                                        and word_table_rows[-1] is not None
                                    ):
                                        word_table_rows[-1] += 1
                                        if word_table_rows[-1] > MAX_EXCEL_COLUMNS:
                                            raise _too_complex(
                                                "wordTableColumns",
                                                word_table_rows[-1],
                                                MAX_EXCEL_COLUMNS,
                                            )
                                continue
                            text_characters += len(element.text or "") + len(element.tail or "")
                            if text_characters > MAX_OFFICE_XML_TEXT_CHARS:
                                raise _too_complex(
                                    "officeXmlTextCharacters",
                                    text_characters,
                                    MAX_OFFICE_XML_TEXT_CHARS,
                                )
                            if lowered_name == "word/document.xml":
                                if local_name == "tbl":
                                    word_tables += 1
                                    if word_table_grids:
                                        word_table_grids.pop()
                                    if word_table_rows:
                                        word_table_rows.pop()
                                    if word_tables > MAX_WORD_TABLES:
                                        raise _too_complex(
                                            "wordTables", word_tables, MAX_WORD_TABLES
                                        )
                                elif local_name == "tc":
                                    word_cells += 1
                                    if word_cells > MAX_WORD_TABLE_CELLS:
                                        raise _too_complex(
                                            "wordTableCells",
                                            word_cells,
                                            MAX_WORD_TABLE_CELLS,
                                        )
                                elif local_name == "p":
                                    word_paragraphs += 1
                                    if word_paragraphs > MAX_WORD_PARAGRAPHS:
                                        raise _too_complex(
                                            "wordParagraphs",
                                            word_paragraphs,
                                            MAX_WORD_PARAGRAPHS,
                                        )
                                elif (
                                    local_name == "gridSpan"
                                    and word_table_rows
                                    and word_table_rows[-1] is not None
                                ):
                                    try:
                                        span = int(element.attrib[qn("w:val")])
                                        if span < 1:
                                            raise ValueError
                                    except (KeyError, TypeError, ValueError) as exc:
                                        raise JobFailure(
                                            ErrorCode.OFFICE_CONVERSION_FAILED,
                                            "The Word document contains an invalid table span",
                                        ) from exc
                                    word_table_rows[-1] += span - 1
                                    if word_table_rows[-1] > MAX_EXCEL_COLUMNS:
                                        raise _too_complex(
                                            "wordTableColumns",
                                            word_table_rows[-1],
                                            MAX_EXCEL_COLUMNS,
                                        )
                                elif local_name == "tr" and word_table_rows:
                                    word_table_rows[-1] = None
                            elif lowered_name == "xl/workbook.xml" and local_name == "sheet":
                                spreadsheet_sheets += 1
                                if spreadsheet_sheets > MAX_SPREADSHEET_SHEETS:
                                    raise _too_complex(
                                        "spreadsheetSheets",
                                        spreadsheet_sheets,
                                        MAX_SPREADSHEET_SHEETS,
                                    )
                            elif lowered_name.startswith("xl/worksheets/"):
                                if local_name == "mergeCell":
                                    reference = element.attrib.get("ref")
                                    try:
                                        min_col, min_row, max_col, max_row = range_boundaries(
                                            reference
                                        )
                                        if (
                                            None in {min_col, min_row, max_col, max_row}
                                            or min_col < 1
                                            or min_row < 1
                                            or max_col < min_col
                                            or max_row < min_row
                                        ):
                                            raise ValueError
                                    except (TypeError, ValueError) as exc:
                                        raise JobFailure(
                                            ErrorCode.OFFICE_CONVERSION_FAILED,
                                            "The spreadsheet contains an invalid merged range",
                                        ) from exc
                                    if max_col > MAX_EXCEL_COLUMNS:
                                        raise _too_complex(
                                            "spreadsheetColumns",
                                            max_col,
                                            MAX_EXCEL_COLUMNS,
                                        )
                                    spreadsheet_merged_cells += (max_row - min_row + 1) * (
                                        max_col - min_col + 1
                                    )
                                    if spreadsheet_merged_cells > MAX_OUTPUT_TABLE_CELLS:
                                        raise _too_complex(
                                            "spreadsheetMergedCells",
                                            spreadsheet_merged_cells,
                                            MAX_OUTPUT_TABLE_CELLS,
                                        )
                                elif local_name == "c":
                                    spreadsheet_materialized_cells += 1
                                    if spreadsheet_materialized_cells > MAX_OUTPUT_TABLE_CELLS:
                                        raise _too_complex(
                                            "spreadsheetMaterializedCells",
                                            spreadsheet_materialized_cells,
                                            MAX_OUTPUT_TABLE_CELLS,
                                        )
                                    if any(
                                        child.tag.rsplit("}", 1)[-1] in {"v", "f", "is"}
                                        for child in element
                                    ):
                                        spreadsheet_cells += 1
                                        if spreadsheet_cells > MAX_SPREADSHEET_NONEMPTY_CELLS:
                                            raise _too_complex(
                                                "spreadsheetNonemptyCells",
                                                spreadsheet_cells,
                                                MAX_SPREADSHEET_NONEMPTY_CELLS,
                                            )
                            element.clear()
                            depth -= 1
    except JobFailure:
        raise
    except (DefusedXmlException, DefusedET.ParseError, OSError, zipfile.BadZipFile) as exc:
        raise JobFailure(
            ErrorCode.OFFICE_CONVERSION_FAILED,
            "The Office package contains invalid XML",
        ) from exc


def _normalize_raster(blob: bytes) -> tuple[bytes, int, int] | None:
    if len(blob) > MAX_MEDIA_BYTES:
        raise _too_complex("embeddedMediaBytes", len(blob), MAX_MEDIA_BYTES)
    try:
        with Image.open(io.BytesIO(blob)) as source:
            width, height = source.size
            if width < 1 or height < 1:
                return None
            pixels = width * height
            if pixels > MAX_EMBEDDED_IMAGE_PIXELS:
                raise _too_complex("embeddedImagePixels", pixels, MAX_EMBEDDED_IMAGE_PIXELS)
            if source.format in {"JPEG", "PNG"}:
                source.verify()
                return blob, width, height
            source.load()
            converted = io.BytesIO()
            if source.mode not in {"RGB", "RGBA", "L", "LA"}:
                source = source.convert("RGBA" if "transparency" in source.info else "RGB")
            source.save(converted, "PNG")
            return converted.getvalue(), width, height
    except (OSError, UnidentifiedImageError):
        return None


def _libreoffice_convert(
    source: Path,
    output_dir: Path,
    extension: str,
    deadline: float | None = None,
) -> tuple[Path, Path]:
    _check_deadline(deadline)
    working = output_dir / f".office-convert-{uuid.uuid4().hex}"
    profile = working / "profile"
    converted_dir = working / "result"
    profile.mkdir(parents=True)
    profile_user = profile / "user"
    profile_user.mkdir()
    (profile_user / "registrymodifications.xcu").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
        '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
        "</item>"
        '<item oor:path="/org.openoffice.Office.Writer/Content/Update">'
        '<prop oor:name="Link" oor:op="fuse"><value>1</value></prop>'
        "</item>"
        '<item oor:path="/org.openoffice.Office.Calc/Content/Update">'
        '<prop oor:name="Link" oor:op="fuse"><value>1</value></prop>'
        "</item>"
        '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
        '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>1</value></prop>'
        '<prop oor:name="ODFRecalcMode" oor:op="fuse"><value>1</value></prop>'
        "</item></oor:items>",
        encoding="utf-8",
    )
    converted_dir.mkdir()
    filter_name = "Office Open XML Text" if extension == ".docx" else "Calc MS Excel 2007 XML"
    command = [
        shutil.which("libreoffice") or shutil.which("soffice") or "libreoffice",
        "--headless",
        "--invisible",
        "--nologo",
        "--nodefault",
        "--nofirststartwizard",
        "--nocrashreport",
        "--nolockcheck",
        "--norestore",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to",
        f"{extension.removeprefix('.')}:{filter_name}",
        "--outdir",
        str(converted_dir),
        str(source),
    ]
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=_remaining_timeout(deadline),
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        shutil.rmtree(working, ignore_errors=True)
        raise JobFailure(ErrorCode.TIMEOUT, "Office conversion timed out") from exc
    generated = next(converted_dir.glob(f"*{extension}"), None)
    diagnostic = f"{completed.stdout}\n{completed.stderr}".lower()
    if completed.returncode != 0 or generated is None:
        shutil.rmtree(working, ignore_errors=True)
        if "password" in diagnostic or "encrypted" in diagnostic:
            raise JobFailure(
                ErrorCode.OFFICE_PASSWORD_PROTECTED,
                "Password-protected Office files are not supported",
            )
        raise JobFailure(
            ErrorCode.OFFICE_CONVERSION_FAILED,
            "LibreOffice could not convert the legacy Office file",
        )
    _check_deadline(deadline)
    return generated, working


def _paragraph_style(
    paragraph: Paragraph,
) -> tuple[bool, bool, str | None, str | None, str | None, float | None, str | None]:
    text_runs = [run for run in paragraph.runs if run.text]
    style_name = paragraph.style.name.casefold() if paragraph.style and paragraph.style.name else ""
    bold = style_name.startswith(("heading", "title", "заголов")) or bool(
        text_runs and all(run.bold for run in text_runs)
    )
    italic = bool(text_runs and all(run.italic for run in text_runs))
    alignment = _paragraph_alignment(paragraph)
    first_run = text_runs[0] if text_runs else None
    font_name = first_run.font.name if first_run is not None else None
    font_size = first_run.font.size.pt if first_run is not None and first_run.font.size else None
    font_color = None
    if first_run is not None and first_run.font.color and first_run.font.color.rgb:
        font_color = str(first_run.font.color.rgb)[-6:]
    hyperlink = _safe_hyperlink(next((link.url for link in paragraph.hyperlinks if link.url), None))
    return bold, italic, alignment, hyperlink, font_name, font_size, font_color


def _paragraph_alignment(paragraph: Paragraph) -> str | None:
    """Read both strict and transitional OOXML alignment values safely."""
    nodes = paragraph._p.xpath("./w:pPr/w:jc")
    raw_value = nodes[0].get(qn("w:val")) if nodes else None
    if raw_value:
        normalized = raw_value.casefold()
        if normalized in {"start", "left"}:
            return "left"
        if normalized in {"end", "right"}:
            return "right"
        if normalized == "center":
            return "center"
        if normalized in {
            "both",
            "distribute",
            "highkashida",
            "lowkashida",
            "mediumkashida",
            "thaidistribute",
        }:
            return "justify"

    try:
        value = paragraph.alignment
    except ValueError:
        return None
    return {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.DISTRIBUTE: "justify",
    }.get(value)


def _paragraph_cell(paragraph: Paragraph) -> CellData:
    value = paragraph.text
    paragraph_properties = paragraph._p.pPr
    style_name = paragraph.style.name.casefold() if paragraph.style and paragraph.style.name else ""
    direct_numbering = paragraph_properties is not None and paragraph_properties.numPr is not None
    if value and (direct_numbering or style_name.startswith("list bullet")):
        value = f"• {value}"
    elif value and style_name.startswith("list number"):
        value = f"1. {value}"
    bold, italic, alignment, hyperlink, font_name, font_size, font_color = _paragraph_style(
        paragraph
    )
    return CellData(
        value=value,
        bold=bold,
        italic=italic,
        font_name=font_name,
        font_size=font_size,
        font_color=font_color,
        alignment=alignment,
        hyperlink=hyperlink,
    )


def _paragraph_rasters(
    paragraph: Paragraph,
    *,
    row: int,
    column: int,
    budget: MediaBudget,
) -> tuple[list[RasterPlacement], bool]:
    placements: list[RasterPlacement] = []
    unsupported = False
    for blip in paragraph._p.xpath(".//a:blip"):
        relationship_id = blip.get(qn("r:embed"))
        if not relationship_id:
            unsupported = True
            continue
        part = paragraph.part.related_parts.get(relationship_id)
        if part is None or part.content_type not in RASTER_CONTENT_TYPES:
            unsupported = True
            continue
        normalized = _normalize_raster(part.blob)
        if normalized is None:
            unsupported = True
            continue
        blob, width, height = normalized
        budget.add(blob, width, height)
        placements.append(RasterPlacement(blob, row, column + len(placements), width, height))
    drawings = paragraph._p.xpath(".//w:drawing | .//w:pict | .//w:object")
    if drawings and not paragraph._p.xpath(".//a:blip"):
        unsupported = True
    return placements, unsupported


def _cell_style(paragraphs: Iterable[Paragraph], cell_element: Any) -> CellData:
    paragraphs = list(paragraphs)
    value = "\n".join(paragraph.text for paragraph in paragraphs)
    paragraph = paragraphs[0] if paragraphs else None
    styled = _paragraph_cell(paragraph) if paragraph is not None else CellData()
    styled.value = value
    shading = cell_element.tcPr.find(qn("w:shd")) if cell_element.tcPr is not None else None
    if shading is not None:
        fill = shading.get(qn("w:fill"))
        if fill and fill.casefold() not in {"auto", "none"}:
            styled.fill = fill[-6:]
    styled.border = True
    return styled


def _word_table_data(
    table: Table,
    *,
    first_row: int,
    locale: str,
    budget: MediaBudget,
) -> tuple[
    list[list[CellData]],
    list[tuple[int, int, int, int]],
    list[RasterPlacement],
    bool,
    list[tuple[str, Table, CellData]],
    dict[int, float],
]:
    rows: list[list[CellData]] = []
    images: list[RasterPlacement] = []
    unsupported = False
    nested_tables: list[tuple[str, Table, CellData]] = []
    cell_origins: dict[Any, tuple[int, int]] = {}
    merge_bounds: dict[Any, list[int]] = {}
    for row_index, row in enumerate(table.rows, 1):
        result_row: list[CellData] = []
        for column_index, cell in enumerate(row.cells, 1):
            identity = cell._tc
            if identity not in cell_origins:
                cell_origins[identity] = (row_index, column_index)
                merge_bounds[identity] = [row_index, column_index, row_index, column_index]
                styled_cell = _cell_style(cell.paragraphs, cell._tc)
                result_row.append(styled_cell)
                for nested_number, nested in enumerate(cell.tables, 1):
                    reference = f"R{row_index}C{column_index}.{nested_number}"
                    reference_text = (
                        f"Вложенная таблица: {reference}"
                        if locale == "ru"
                        else f"Nested table: {reference}"
                    )
                    styled_cell.value = (
                        f"{styled_cell.value}\n↳ {reference_text}"
                        if styled_cell.value
                        else f"↳ {reference_text}"
                    )
                    nested_tables.append((reference, nested, styled_cell))
                for paragraph in cell.paragraphs:
                    raster, skipped = _paragraph_rasters(
                        paragraph,
                        row=first_row + row_index - 1,
                        column=column_index,
                        budget=budget,
                    )
                    images.extend(raster)
                    unsupported = unsupported or skipped
            else:
                bounds = merge_bounds[identity]
                bounds[2] = max(bounds[2], row_index)
                bounds[3] = max(bounds[3], column_index)
                result_row.append(CellData())
        rows.append(result_row)
    merges = [
        (first_row + top - 1, left, first_row + bottom - 1, right)
        for top, left, bottom, right in merge_bounds.values()
        if top != bottom or left != right
    ]
    widths: dict[int, float] = {}
    grid = table._tbl.tblGrid
    if grid is not None:
        for column, grid_column in enumerate(grid.gridCol_lst, 1):
            raw_width = grid_column.get(qn("w:w"))
            if raw_width and raw_width.isdigit():
                widths[column] = min(60.0, max(4.0, int(raw_width) / 140))
    return rows, merges, images, unsupported, nested_tables, widths


def _word_has_unsupported_objects(document: DocumentType) -> bool:
    tags = {element.tag.rsplit("}", 1)[-1] for element in document.element.iter()}
    if tags & {
        "altChunk",
        "chart",
        "cxnSp",
        "graphicFrame",
        "grpSp",
        "imagedata",
        "object",
        "oleObject",
        "oMath",
        "relIds",
        "shape",
        "sp",
        "textbox",
        "txbxContent",
        "wgp",
        "wsp",
    }:
        return True
    if any(
        link.url and _safe_hyperlink(link.url) is None
        for paragraph in document.paragraphs
        for link in paragraph.hyperlinks
    ):
        return True
    if any(
        relationship.reltype == RT.HYPERLINK and _safe_hyperlink(relationship.target_ref) is None
        for relationship in document.part.rels.values()
    ):
        return True
    if any(
        _word_related_part_has_unmapped_content(relationship)
        for relationship in document.part.rels.values()
    ):
        return True
    supported_relationships = {
        RT.COMMENTS,
        RT.CORE_PROPERTIES,
        RT.CUSTOM_PROPERTIES,
        RT.CUSTOM_XML,
        RT.ENDNOTES,
        RT.FONT_TABLE,
        RT.FOOTER,
        RT.FOOTNOTES,
        RT.HEADER,
        RT.HYPERLINK,
        RT.IMAGE,
        RT.NUMBERING,
        RT.OFFICE_DOCUMENT,
        RT.SETTINGS,
        RT.STYLES,
        RT.THEME,
        RT.WEB_SETTINGS,
        "http://schemas.microsoft.com/office/2007/relationships/stylesWithEffects",
    }
    return any(
        relationship.reltype not in supported_relationships
        for relationship in document.part.rels.values()
    )


def _word_related_part_has_unmapped_content(relationship: Any) -> bool:
    relation_type = relationship.reltype
    lossy_types = {
        RT.COMMENTS,
        RT.CUSTOM_XML,
        RT.ENDNOTES,
        RT.FOOTER,
        RT.FOOTNOTES,
        RT.HEADER,
    }
    if relation_type not in lossy_types:
        return False
    try:
        custom_xml_elements = 0
        custom_xml_root = ""
        with io.BytesIO(relationship.target_part.blob) as source:
            for _event, element in DefusedET.iterparse(source, events=("end",)):
                local_name = element.tag.rsplit("}", 1)[-1]
                if relation_type == RT.CUSTOM_XML:
                    custom_xml_elements += 1
                    custom_xml_root = local_name
                    if custom_xml_elements > 1 or (element.text or "").strip():
                        return True
                if relation_type in {RT.COMMENTS} and local_name == "comment":
                    return True
                if relation_type in {RT.FOOTNOTES, RT.ENDNOTES} and local_name in {
                    "footnote",
                    "endnote",
                }:
                    raw_id = next(
                        (
                            value
                            for name, value in element.attrib.items()
                            if name.rsplit("}", 1)[-1] == "id"
                        ),
                        "",
                    )
                    if raw_id not in {"-1", "0"}:
                        return True
                if relation_type in {RT.HEADER, RT.FOOTER}:
                    if local_name in {
                        "drawing",
                        "fldSimple",
                        "imagedata",
                        "instrText",
                        "object",
                        "pict",
                        "shape",
                        "tbl",
                        "textbox",
                        "txbxContent",
                    }:
                        return True
                    if local_name == "t" and (element.text or "").strip():
                        return True
                element.clear()
        if relation_type == RT.CUSTOM_XML:
            return custom_xml_root != "Sources"
    except (
        AttributeError,
        DefusedXmlException,
        DefusedET.ParseError,
        KeyError,
        OSError,
        ValueError,
    ):
        return True
    return False


def _parse_word(path: Path, locale: str, deadline: float | None = None) -> ParsedOffice:
    _safe_xml_package(path, deadline)
    _check_deadline(deadline)
    document = Document(path)
    _check_deadline(deadline)
    rows: list[list[CellData]] = []
    merges: list[tuple[int, int, int, int]] = []
    images: list[RasterPlacement] = []
    table_count = 0
    table_cells = 0
    media_budget = MediaBudget()
    maximum_width = 1
    column_widths: dict[int, float] = {}
    paragraph_rows: list[int] = []
    unsupported = _word_has_unsupported_objects(document)

    def separate_block() -> None:
        if rows:
            rows.append([CellData()])

    def append_paragraph(paragraph: Paragraph) -> None:
        nonlocal unsupported
        separate_block()
        row_number = len(rows) + 1
        cell = _paragraph_cell(paragraph)
        raster, skipped = _paragraph_rasters(
            paragraph,
            row=row_number,
            column=1,
            budget=media_budget,
        )
        rows.append([cell])
        paragraph_rows.append(row_number)
        images.extend(raster)
        unsupported = unsupported or skipped

    def append_table(
        table: Table,
        nested_label: str | None = None,
        parent_cell: CellData | None = None,
    ) -> None:
        nonlocal table_count, table_cells, unsupported, maximum_width
        separate_block()
        if nested_label:
            label = "Вложенная таблица" if locale == "ru" else "Nested table"
            rows.append([CellData(f"{label} ({nested_label})", bold=True)])
            paragraph_rows.append(len(rows))
            if parent_cell is not None:
                parent_cell.hyperlink = (
                    f"#'{'Документ' if locale == 'ru' else 'Document'}'!A{len(rows)}"
                )
        table_count += 1
        if table_count > MAX_WORD_TABLES:
            raise _too_complex("wordTables", table_count, MAX_WORD_TABLES)
        row_widths = [len(row.cells) for row in table.rows]
        width = max(row_widths, default=0)
        if width > MAX_EXCEL_COLUMNS:
            raise _too_complex("wordTableColumns", width, MAX_EXCEL_COLUMNS)
        count = sum(row_widths)
        table_cells += count
        if table_cells > MAX_WORD_TABLE_CELLS:
            raise _too_complex("wordTableCells", table_cells, MAX_WORD_TABLE_CELLS)
        (
            table_rows,
            table_merges,
            table_images,
            skipped,
            nested_tables,
            widths,
        ) = _word_table_data(
            table,
            first_row=len(rows) + 1,
            locale=locale,
            budget=media_budget,
        )
        maximum_width = max(maximum_width, max((len(row) for row in table_rows), default=1))
        for column, width in widths.items():
            column_widths[column] = max(column_widths.get(column, 0), width)
        rows.extend(table_rows)
        merges.extend(table_merges)
        images.extend(table_images)
        unsupported = unsupported or skipped
        for reference, nested, reference_cell in nested_tables:
            append_table(nested, reference, reference_cell)

    for block in document.iter_inner_content():
        _check_deadline(deadline)
        if isinstance(block, Paragraph):
            append_paragraph(block)
        elif isinstance(block, Table):
            append_table(block)
    if maximum_width > 1:
        merges.extend((row, 1, row, maximum_width) for row in paragraph_rows)
    return ParsedOffice(
        sheets=[
            SheetData(
                "Документ" if locale == "ru" else "Document",
                rows or [[CellData()]],
                merges,
                images,
                source_columns=list(range(1, maximum_width + 1)),
                column_widths=column_widths,
            )
        ],
        unsupported_objects=unsupported,
    )


def _safe_sheet_title(value: str, used: set[str], number: int) -> str:
    base = re.sub(r"[\[\]:*?/\\]+", "_", value).strip(" '")[:31] or f"Sheet {number}"
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        ending = f" {suffix}"
        candidate = f"{base[: 31 - len(ending)]}{ending}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _write_excel(
    destination: Path,
    parsed: ParsedOffice,
    deadline: float | None = None,
) -> None:
    for sheet in parsed.sheets:
        for row in sheet.rows:
            if len(row) > MAX_EXCEL_COLUMNS:
                raise _too_complex("excelColumns", len(row), MAX_EXCEL_COLUMNS)
            for value in row:
                if len(value.value) > MAX_EXCEL_CELL_CHARACTERS:
                    raise _too_complex(
                        "excelCellCharacters",
                        len(value.value),
                        MAX_EXCEL_CELL_CHARACTERS,
                    )
        if any(right > MAX_EXCEL_COLUMNS for _top, _left, _bottom, right in sheet.merges):
            raise _too_complex(
                "excelColumns",
                max(right for _top, _left, _bottom, right in sheet.merges),
                MAX_EXCEL_COLUMNS,
            )
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for sheet_number, sheet in enumerate(parsed.sheets, 1):
        _check_deadline(deadline)
        worksheet = workbook.create_sheet(_safe_sheet_title(sheet.name, used, sheet_number))
        for row_index, row in enumerate(sheet.rows, 1):
            for column_index, value in enumerate(row, 1):
                cell = worksheet.cell(row_index, column_index)
                cell.value = value.value
                if value.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
                cell.font = Font(
                    name=value.font_name,
                    size=value.font_size,
                    bold=value.bold,
                    italic=value.italic,
                    color=value.font_color,
                )
                if value.fill:
                    cell.fill = PatternFill("solid", fgColor=value.fill)
                if value.alignment:
                    cell.alignment = Alignment(horizontal=value.alignment, wrap_text=True)
                else:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if value.border:
                    side = Side(style="thin", color="B7B7B7")
                    cell.border = Border(left=side, right=side, top=side, bottom=side)
                if value.hyperlink:
                    if value.hyperlink.startswith("#"):
                        cell.hyperlink = Hyperlink(
                            ref=cell.coordinate,
                            location=value.hyperlink[1:],
                        )
                    else:
                        cell.hyperlink = value.hyperlink
                    cell.style = "Hyperlink"
                if value.comment:
                    cell.comment = Comment(value.comment[0], value.comment[1])
                if value.value:
                    worksheet.column_dimensions[cell.column_letter].width = min(
                        60,
                        max(
                            worksheet.column_dimensions[cell.column_letter].width or 8,
                            max((len(line) for line in value.value.splitlines()), default=0) + 2,
                        ),
                    )
        for top, left, bottom, right in sheet.merges:
            if top != bottom or left != right:
                worksheet.merge_cells(
                    start_row=top,
                    start_column=left,
                    end_row=bottom,
                    end_column=right,
                )
        for column, width in sheet.column_widths.items():
            worksheet.column_dimensions[get_column_letter(column)].width = min(60, max(4, width))
        for placement in sheet.images:
            image = ExcelImage(io.BytesIO(placement.data))
            scale = min(1.0, 700 / placement.width, 500 / placement.height)
            image.width = max(1, round(placement.width * scale))
            image.height = max(1, round(placement.height * scale))
            anchor = worksheet.cell(placement.row, placement.column).coordinate
            worksheet.add_image(image, anchor)
            worksheet.row_dimensions[placement.row].height = max(
                worksheet.row_dimensions[placement.row].height or 15,
                image.height * 0.75,
            )
    try:
        workbook.save(destination)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()


DELIMITERS = {"comma": ",", "semicolon": ";", "tab": "\t", "pipe": "|"}


def _looks_numeric(value: str) -> bool:
    try:
        float(value.strip().replace(",", "."))
        return True
    except ValueError:
        return False


def _decode_csv(raw: bytes, requested: str) -> tuple[str, str, bool]:
    if requested not in {"auto", "utf-8", "windows-1251"}:
        raise JobFailure(ErrorCode.CSV_PARSE_FAILED, "Unsupported CSV encoding")
    guessed = requested == "auto"
    if requested != "auto":
        label, codec = (
            ("utf-8", "utf-8-sig") if requested == "utf-8" else ("windows-1251", "cp1251")
        )
        try:
            value = raw.decode(codec, errors="strict")
            if "\x00" not in value:
                return value, label, guessed
        except UnicodeDecodeError as exc:
            raise JobFailure(
                ErrorCode.CSV_PARSE_FAILED,
                "The CSV encoding could not be decoded",
            ) from exc
        raise JobFailure(ErrorCode.CSV_PARSE_FAILED, "The CSV contains binary data")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16", errors="strict"), "utf-16", True
        except UnicodeDecodeError as exc:
            raise JobFailure(ErrorCode.CSV_PARSE_FAILED, "Invalid UTF-16 CSV") from exc
    try:
        value = raw.decode("utf-8-sig", errors="strict")
        if "\x00" not in value:
            return value, "utf-8", True
    except UnicodeDecodeError:
        pass
    matches = list(from_bytes(raw))
    best = matches[0] if matches else None
    normalized_encoding = (
        (best.encoding or "").replace("-", "").replace("_", "").casefold() if best else ""
    )
    competing = matches[1] if len(matches) > 1 else None
    ambiguous = bool(competing is not None and competing.percent_chaos - best.percent_chaos < 5)
    if (
        best is None
        or normalized_encoding not in {"cp1251", "windows1251"}
        or best.percent_chaos > 10
        or ambiguous
    ):
        raise JobFailure(
            ErrorCode.CSV_PARSE_FAILED,
            "CSV encoding detection was not confident; choose an encoding manually",
        )
    value = str(best)
    if "\x00" in value:
        raise JobFailure(ErrorCode.CSV_PARSE_FAILED, "The CSV contains binary data")
    return value, "windows-1251", True


def _parse_csv(
    path: Path,
    options: dict[str, Any],
    locale: str,
    deadline: float | None = None,
) -> tuple[ParsedOffice, dict[str, Any]]:
    _check_deadline(deadline)
    text, encoding, encoding_guessed = _decode_csv(
        path.read_bytes(), str(options.get("csvEncoding", "auto")).lower()
    )
    delimiter_option = str(options.get("csvDelimiter", "auto")).lower()
    delimiter_guessed = delimiter_option == "auto"
    if delimiter_option == "auto":
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        except csv.Error as exc:
            raise JobFailure(
                ErrorCode.CSV_PARSE_FAILED,
                "CSV delimiter detection was not confident; choose a delimiter manually",
            ) from exc
    elif delimiter_option in DELIMITERS:
        delimiter = DELIMITERS[delimiter_option]
    else:
        raise JobFailure(ErrorCode.CSV_PARSE_FAILED, "Unsupported CSV delimiter")
    for physical_line in io.StringIO(text):
        _check_deadline(deadline)
        estimated_columns = physical_line.count(delimiter) + 1
        if estimated_columns > MAX_OUTPUT_TABLE_CELLS:
            raise _too_complex(
                "csvRowCells",
                estimated_columns,
                MAX_OUTPUT_TABLE_CELLS,
            )
    raw_rows: list[list[str]] = []
    nonempty = 0
    materialized_cells = 0
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True):
            _check_deadline(deadline)
            if not any(value for value in row):
                continue
            materialized_cells += len(row)
            if materialized_cells > MAX_OUTPUT_TABLE_CELLS:
                raise _too_complex(
                    "csvMaterializedCells",
                    materialized_cells,
                    MAX_OUTPUT_TABLE_CELLS,
                )
            nonempty += sum(bool(value) for value in row)
            if nonempty > MAX_SPREADSHEET_NONEMPTY_CELLS:
                raise _too_complex(
                    "spreadsheetNonemptyCells",
                    nonempty,
                    MAX_SPREADSHEET_NONEMPTY_CELLS,
                )
            raw_rows.append(row)
    except (csv.Error, UnicodeError) as exc:
        raise JobFailure(ErrorCode.CSV_PARSE_FAILED, "The CSV file could not be parsed") from exc
    widths = {len(row) for row in raw_rows}
    if len(widths) > 1:
        raise JobFailure(
            ErrorCode.CSV_PARSE_FAILED,
            "CSV rows contain an inconsistent number of columns",
            {"widths": sorted(widths)},
        )
    first_row_is_text = bool(raw_rows) and all(
        value.strip() and not _looks_numeric(value) for value in raw_rows[0]
    )
    try:
        has_header = first_row_is_text or (
            bool(raw_rows) and csv.Sniffer().has_header(text[:65536])
        )
    except csv.Error:
        has_header = first_row_is_text
    rows = [
        [
            CellData(
                value=value,
                bold=has_header and row_index == 0,
                fill="D9EAF7" if has_header and row_index == 0 else None,
                border=True,
            )
            for value in row
        ]
        for row_index, row in enumerate(raw_rows)
    ]
    column_count = len(raw_rows[0]) if raw_rows else 1
    parsed = ParsedOffice(
        [
            SheetData(
                _clean_stem(path.name),
                rows or [[CellData()]],
                source_columns=list(range(1, column_count + 1)),
            )
        ]
    )
    details = {
        "delimiter": {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}[delimiter],
        "encoding": encoding,
        "delimiterDetected": delimiter_guessed,
        "encodingDetected": encoding_guessed,
    }
    return parsed, details


def _cell_rgb(cell: Cell) -> str | None:
    color = cell.fill.fgColor
    if cell.fill.fill_type and color.type == "rgb" and color.rgb:
        return color.rgb[-6:]
    return None


def _font_rgb(cell: Cell) -> str | None:
    color = cell.font.color
    if color and color.type == "rgb" and color.rgb:
        return color.rgb[-6:]
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _excel_display_text(value: Any, number_format: str | None) -> str:
    if value is None or isinstance(value, bool) or not number_format:
        return _cell_text(value)
    if isinstance(value, datetime):
        lowered_format = number_format.casefold()
        has_date = any(token in lowered_format for token in ("y", "d"))
        has_time = any(token in lowered_format for token in ("h", "s"))
        if has_date and not has_time:
            return value.date().isoformat()
        if has_time and not has_date:
            return value.time().isoformat(timespec="seconds")
        if value.microsecond:
            return value.isoformat(sep=" ")
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, datetime_time):
        return value.isoformat(timespec="seconds")
    if not isinstance(value, int | float):
        return _cell_text(value)
    normalized = re.sub(r'"[^"]*"|\\.', "", number_format.split(";", 1)[0])
    if is_date_format(number_format):
        return _cell_text(value)
    if "%" in normalized:
        decimal_match = re.search(r"\.(0+|#+)\s*%", normalized)
        decimals = len(decimal_match.group(1)) if decimal_match else 0
        return f"{value * 100:.{decimals}f}%"
    decimal_match = re.search(r"\.(0+|#+)", normalized)
    decimals = len(decimal_match.group(1)) if decimal_match else 0
    use_grouping = "," in normalized
    numeric = f"{value:,.{decimals}f}" if use_grouping else f"{value:.{decimals}f}"
    currency_match = re.search(r"[$₽€£¥]", normalized)
    if currency_match:
        return f"{currency_match.group(0)}{numeric}"
    if use_grouping or decimal_match:
        return numeric
    return _cell_text(value)


def _excel_cell_data(source: Cell, cached: Cell) -> tuple[CellData, bool, bool]:
    formula = source.data_type == "f"
    value = cached.value if formula and cached.value is not None else source.value
    unsupported_formula = isinstance(source.value, DataTableFormula)
    if formula and cached.value is None and not isinstance(value, str):
        value = getattr(value, "text", "")
    hyperlink = None
    unsafe_hyperlink = False
    if source.hyperlink:
        raw_hyperlink = source.hyperlink.target or source.hyperlink.location
        hyperlink = _safe_hyperlink(raw_hyperlink)
        unsafe_hyperlink = bool(raw_hyperlink and hyperlink is None)
    comment = (source.comment.text, source.comment.author or "") if source.comment else None
    return (
        CellData(
            value=_excel_display_text(value, source.number_format),
            bold=bool(source.font.bold),
            italic=bool(source.font.italic),
            font_name=source.font.name,
            font_size=float(source.font.sz) if source.font.sz else None,
            font_color=_font_rgb(source),
            fill=_cell_rgb(source),
            alignment=source.alignment.horizontal,
            border=any(
                side.style
                for side in (
                    source.border.left,
                    source.border.right,
                    source.border.top,
                    source.border.bottom,
                )
            ),
            hyperlink=hyperlink,
            comment=comment,
        ),
        formula,
        unsafe_hyperlink or unsupported_formula,
    )


def _excel_image(
    image: Any,
    row: int,
    column: int,
    budget: MediaBudget,
) -> RasterPlacement | None:
    try:
        blob = image._data()
    except (OSError, ValueError):
        return None
    normalized = _normalize_raster(blob)
    if normalized is None:
        return None
    data, width, height = normalized
    budget.add(data, width, height)
    return RasterPlacement(data, row, column, width, height)


def _excel_package_objects(path: Path, deadline: float | None) -> bool:
    unsupported = False
    try:
        with zipfile.ZipFile(path) as package:
            for info in package.infolist():
                _check_deadline(deadline)
                lowered = info.filename.casefold()
                if lowered.startswith(
                    (
                        "xl/activex/",
                        "xl/ctrlprops/",
                        "xl/diagrams/",
                        "xl/embeddings/",
                        "xl/externallinks/",
                        "xl/persons/",
                        "xl/pivotcache/",
                        "xl/pivottables/",
                        "xl/slicers/",
                        "xl/threadedcomments/",
                    )
                ):
                    unsupported = True
                if lowered.startswith("xl/drawings/") and lowered.endswith(".xml"):
                    with package.open(info) as source:
                        for _event, element in DefusedET.iterparse(source, events=("end",)):
                            if element.tag.rsplit("}", 1)[-1] in {
                                "cxnSp",
                                "graphicFrame",
                                "grpSp",
                                "sp",
                            }:
                                unsupported = True
                            element.clear()
    except JobFailure:
        raise
    except (DefusedXmlException, DefusedET.ParseError, OSError, zipfile.BadZipFile) as exc:
        raise JobFailure(
            ErrorCode.OFFICE_CONVERSION_FAILED,
            "The spreadsheet package contains invalid drawing data",
        ) from exc
    return unsupported


def _parse_excel(path: Path, deadline: float | None = None) -> ParsedOffice:
    _safe_xml_package(path, deadline)
    _check_deadline(deadline)
    package_unsupported = _excel_package_objects(path, deadline)
    formula_book = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    _check_deadline(deadline)
    value_book = load_workbook(path, data_only=True, read_only=False, keep_links=False)
    _check_deadline(deadline)
    try:
        if len(formula_book.worksheets) > MAX_SPREADSHEET_SHEETS:
            raise _too_complex(
                "spreadsheetSheets",
                len(formula_book.worksheets),
                MAX_SPREADSHEET_SHEETS,
            )
        sheets: list[SheetData] = []
        nonempty = 0
        formulas = False
        unsupported = package_unsupported or bool(getattr(formula_book, "_external_links", []))
        media_budget = MediaBudget()
        for worksheet in formula_book.worksheets:
            _check_deadline(deadline)
            cached_sheet = value_book[worksheet.title]
            unsupported = unsupported or bool(
                worksheet.data_validations.count
                or len(worksheet.conditional_formatting)
                or worksheet._pivots
                or worksheet._charts
            )
            stored_cells = list(worksheet._cells.values())
            meaningful_cells = [
                cell
                for cell in stored_cells
                if cell.value is not None or cell.comment is not None or cell.hyperlink is not None
            ]
            if not meaningful_cells and not worksheet._images:
                continue
            used_rows = {cell.row for cell in meaningful_cells}
            used_columns = {cell.column for cell in meaningful_cells}
            for merged in worksheet.merged_cells.ranges:
                merged_cells = (merged.max_row - merged.min_row + 1) * (
                    merged.max_col - merged.min_col + 1
                )
                if merged_cells > MAX_OUTPUT_TABLE_CELLS:
                    raise _too_complex(
                        "spreadsheetMergedCells",
                        merged_cells,
                        MAX_OUTPUT_TABLE_CELLS,
                    )
                used_rows.update(range(merged.min_row, merged.max_row + 1))
                used_columns.update(range(merged.min_col, merged.max_col + 1))
                projected_cells = len(used_rows) * max(1, len(used_columns))
                if projected_cells > MAX_OUTPUT_TABLE_CELLS:
                    raise _too_complex(
                        "outputTableCells",
                        projected_cells,
                        MAX_OUTPUT_TABLE_CELLS,
                    )
            for image in worksheet._images:
                anchor = getattr(image.anchor, "_from", None)
                if anchor is not None:
                    used_rows.add(anchor.row + 1)
                    used_columns.add(anchor.col + 1)
            sorted_rows = sorted(used_rows)
            sorted_columns = sorted(used_columns)
            if len(sorted_rows) * max(1, len(sorted_columns)) > MAX_OUTPUT_TABLE_CELLS:
                raise _too_complex(
                    "outputTableCells",
                    len(sorted_rows) * len(sorted_columns),
                    MAX_OUTPUT_TABLE_CELLS,
                )
            row_map = {original: index for index, original in enumerate(sorted_rows, 1)}
            column_map = {original: index for index, original in enumerate(sorted_columns, 1)}
            rows: list[list[CellData]] = []
            for original_row in sorted_rows:
                _check_deadline(deadline)
                result_row: list[CellData] = []
                for original_column in sorted_columns:
                    source = worksheet.cell(original_row, original_column)
                    cached = cached_sheet.cell(original_row, original_column)
                    data, is_formula, unsafe_link = _excel_cell_data(source, cached)
                    formulas = formulas or is_formula
                    unsupported = unsupported or unsafe_link
                    if source.value is not None:
                        nonempty += 1
                        if nonempty > MAX_SPREADSHEET_NONEMPTY_CELLS:
                            raise _too_complex(
                                "spreadsheetNonemptyCells",
                                nonempty,
                                MAX_SPREADSHEET_NONEMPTY_CELLS,
                            )
                    result_row.append(data)
                rows.append(result_row)
            merges: list[tuple[int, int, int, int]] = []
            for merged in worksheet.merged_cells.ranges:
                if all(
                    coordinate in mapping
                    for coordinate, mapping in (
                        (merged.min_row, row_map),
                        (merged.max_row, row_map),
                        (merged.min_col, column_map),
                        (merged.max_col, column_map),
                    )
                ):
                    merges.append(
                        (
                            row_map[merged.min_row],
                            column_map[merged.min_col],
                            row_map[merged.max_row],
                            column_map[merged.max_col],
                        )
                    )
            images: list[RasterPlacement] = []
            for image in worksheet._images:
                anchor = getattr(image.anchor, "_from", None)
                if anchor is None:
                    unsupported = True
                    continue
                placement = _excel_image(
                    image,
                    row_map.get(anchor.row + 1, 1),
                    column_map.get(anchor.col + 1, 1),
                    media_budget,
                )
                if placement is None:
                    unsupported = True
                    continue
                images.append(placement)
            widths = {
                column_map[original]: float(
                    worksheet.column_dimensions[get_column_letter(original)].width or 8.43
                )
                for original in sorted_columns
            }
            sheets.append(
                SheetData(
                    worksheet.title,
                    rows or [[CellData()]],
                    merges,
                    images,
                    hidden=worksheet.sheet_state != "visible",
                    source_columns=sorted_columns,
                    column_widths=widths,
                )
            )
        return ParsedOffice(sheets, formulas, unsupported)
    finally:
        formula_book.close()
        value_book.close()


def _sheet_bands(sheet: SheetData, locale: str) -> list[tuple[str | None, SheetData]]:
    width = max((len(row) for row in sheet.rows), default=1)
    if width <= MAX_WORD_COLUMNS_PER_BAND:
        return [(None, sheet)]
    source_columns = sheet.source_columns or list(range(1, width + 1))
    bands: list[tuple[str | None, SheetData]] = []
    for start in range(1, width, MAX_WORD_COLUMNS_PER_BAND - 1):
        selected = [0, *range(start, min(width, start + MAX_WORD_COLUMNS_PER_BAND - 1))]
        mapping = {old + 1: new + 1 for new, old in enumerate(selected)}
        rows = [
            [row[index] if index < len(row) else CellData() for index in selected]
            for row in sheet.rows
        ]
        merges = [
            (top, mapping[left], bottom, mapping[right])
            for top, left, bottom, right in sheet.merges
            if left in mapping
            and right in mapping
            and all(column in mapping for column in range(left, right + 1))
        ]
        images = [
            RasterPlacement(
                placement.data,
                placement.row,
                mapping[placement.column],
                placement.width,
                placement.height,
            )
            for placement in sheet.images
            if placement.column in mapping
        ]
        selected_sources = [source_columns[index] for index in selected]
        first_label = get_column_letter(selected_sources[0])
        range_start = get_column_letter(selected_sources[1])
        range_end = get_column_letter(selected_sources[-1])
        columns_label = DOCUMENT_LABELS[locale]["columns"]
        label = (
            f"{columns_label} {first_label}–{range_end}"
            if selected_sources[1] == selected_sources[0] + 1
            else f"{columns_label} {first_label} + {range_start}–{range_end}"
        )
        widths = {
            mapping[old]: width_value
            for old, width_value in sheet.column_widths.items()
            if old in mapping
        }
        bands.append(
            (
                label,
                SheetData(
                    sheet.name,
                    rows,
                    merges,
                    images,
                    sheet.hidden,
                    [source_columns[index] for index in selected],
                    widths,
                ),
            )
        )
    return bands


def _add_hyperlink(paragraph: Paragraph, text: str, url: str) -> Run:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run_element.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run_element.append(text_element)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)
    return Run(run_element, paragraph)


def _apply_word_cell(document: DocumentType, target: Any, value: CellData) -> None:
    paragraph = target.paragraphs[0]
    if value.hyperlink and value.value:
        run = _add_hyperlink(paragraph, value.value, value.hyperlink)
    else:
        run = paragraph.add_run(value.value)
    run.bold = value.bold
    run.italic = value.italic
    run.font.name = value.font_name or "Aptos"
    run.font.size = Pt(max(7.0, value.font_size or 9.0))
    if value.font_color:
        try:
            run.font.color.rgb = RGBColor.from_string(value.font_color)
        except ValueError:
            pass
    if value.fill:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), value.fill)
        target._tc.get_or_add_tcPr().append(shading)
    if value.border:
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "B7B7B7")
            borders.append(border)
        target._tc.get_or_add_tcPr().append(borders)
    paragraph.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(value.alignment)
    if value.comment:
        document.add_comment(run, text=value.comment[0], author=value.comment[1])


def _add_word_image(target: Any, placement: RasterPlacement) -> None:
    paragraph = target.paragraphs[-1]
    run = paragraph.add_run()
    max_width = 6.2 if placement.column == 1 else 2.5
    width = min(max_width, placement.width / 96)
    run.add_picture(io.BytesIO(placement.data), width=Inches(max(0.2, width)))


def _write_word(
    destination: Path,
    parsed: ParsedOffice,
    locale: str,
    deadline: float | None = None,
) -> None:
    document = Document()

    def configure_section(section: Any, wide: bool) -> None:
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.35)
        section.right_margin = Inches(0.35)
        if wide:
            section.orientation = WD_ORIENT.LANDSCAPE
            if section.page_width < section.page_height:
                section.page_width, section.page_height = section.page_height, section.page_width
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            if section.page_width > section.page_height:
                section.page_width, section.page_height = section.page_height, section.page_width

    for sheet_number, sheet in enumerate(parsed.sheets):
        _check_deadline(deadline)
        section = (
            document.sections[0] if sheet_number == 0 else document.add_section(WD_SECTION.NEW_PAGE)
        )
        sheet_is_wide = max((len(row) for row in sheet.rows), default=1) > 6
        configure_section(section, sheet_is_wide)
        hidden_label = DOCUMENT_LABELS[locale]["hidden"]
        sheet_label = f"{sheet.name} [{hidden_label}]" if sheet.hidden else sheet.name
        document.add_heading(sheet_label, level=1)
        for band_label, band in _sheet_bands(sheet, locale):
            if band_label:
                document.add_heading(band_label, level=2)
            width = max((len(row) for row in band.rows), default=1)
            if len(band.rows) * max(width, 1) > MAX_OUTPUT_TABLE_CELLS:
                raise _too_complex(
                    "outputTableCells",
                    len(band.rows) * max(width, 1),
                    MAX_OUTPUT_TABLE_CELLS,
                )
            table = document.add_table(rows=max(1, len(band.rows)), cols=max(1, width))
            table.style = "Table Grid"
            table.autofit = False
            total_width = sum(
                band.column_widths.get(column, 8.43) for column in range(1, width + 1)
            )
            available_inches = 10.2 if sheet_is_wide else 7.1
            for row_index, row in enumerate(band.rows):
                for column_index, value in enumerate(row):
                    target = table.cell(row_index, column_index)
                    _apply_word_cell(document, target, value)
                    relative_width = band.column_widths.get(column_index + 1, 8.43)
                    target.width = Inches(
                        max(0.45, available_inches * relative_width / max(total_width, 1))
                    )
            for top, left, bottom, right in band.merges:
                table.cell(top - 1, left - 1).merge(table.cell(bottom - 1, right - 1))
            for placement in band.images:
                _add_word_image(
                    table.cell(placement.row - 1, placement.column - 1),
                    placement,
                )
    if not parsed.sheets:
        configure_section(document.sections[0], False)
        document.add_paragraph(DOCUMENT_LABELS[locale]["no_nonempty_sheets"])
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    try:
        document.save(destination)
    except Exception:
        destination.unlink(missing_ok=True)
        raise


def _package_outputs(output_dir: Path, outputs: list[Path], name: str) -> Path:
    if len(outputs) == 1:
        return outputs[0]
    archive_path = output_dir / name
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for output in outputs:
            archive.write(output, output.name)
    for output in outputs:
        output.unlink(missing_ok=True)
    return archive_path


def word_to_excel(
    files: list[dict[str, Any]], output_dir: Path, options: dict[str, Any]
) -> OperationResult:
    locale = _locale(options)
    deadline = _conversion_deadline(options)
    outputs: list[Path] = []
    warnings = [_warning("LAYOUT_SIMPLIFIED", locale)]
    try:
        for index, item in enumerate(files, 1):
            _check_deadline(deadline)
            source = Path(item["path"])
            if source.suffix.lower() not in SUPPORTED_WORD_INPUTS:
                raise JobFailure(ErrorCode.OFFICE_CONVERSION_FAILED, "Unsupported Word input")
            parse_path = source
            cleanup: Path | None = None
            if source.suffix.lower() != ".docx":
                parse_path, cleanup = _libreoffice_convert(
                    source,
                    output_dir,
                    ".docx",
                    deadline,
                )
            try:
                parsed = _parse_word(parse_path, locale, deadline)
            finally:
                if cleanup is not None:
                    shutil.rmtree(cleanup, ignore_errors=True)
            if parsed.unsupported_objects:
                warnings.append(_warning("OBJECTS_SKIPPED", locale))
            destination = output_dir / f"{_clean_stem(item['original'])}_{index}.xlsx"
            _write_excel(destination, parsed, deadline)
            outputs.append(destination)
    except JobFailure:
        raise
    except Exception as exc:
        raise JobFailure(
            ErrorCode.OFFICE_CONVERSION_FAILED,
            "The Word file could not be converted",
        ) from exc
    return OperationResult(
        _package_outputs(output_dir, outputs, "word-to-excel.zip"),
        _deduplicate_warnings(warnings),
    )


def excel_to_word(
    files: list[dict[str, Any]], output_dir: Path, options: dict[str, Any]
) -> OperationResult:
    locale = _locale(options)
    deadline = _conversion_deadline(options)
    outputs: list[Path] = []
    warnings = [_warning("LAYOUT_SIMPLIFIED", locale)]
    try:
        for index, item in enumerate(files, 1):
            _check_deadline(deadline)
            source = Path(item["path"])
            if source.suffix.lower() not in SUPPORTED_EXCEL_INPUTS:
                raise JobFailure(ErrorCode.OFFICE_CONVERSION_FAILED, "Unsupported Excel input")
            cleanup: Path | None = None
            csv_details: dict[str, Any] | None = None
            if source.suffix.lower() == ".csv":
                parsed, csv_details = _parse_csv(source, options, locale, deadline)
            else:
                parse_path = source
                if source.suffix.lower() != ".xlsx":
                    parse_path, cleanup = _libreoffice_convert(
                        source,
                        output_dir,
                        ".xlsx",
                        deadline,
                    )
                try:
                    parsed = _parse_excel(parse_path, deadline)
                finally:
                    if cleanup is not None:
                        shutil.rmtree(cleanup, ignore_errors=True)
            if parsed.formulas:
                warnings.append(_warning("FORMULAS_AS_VALUES", locale))
            if parsed.unsupported_objects:
                warnings.append(_warning("OBJECTS_SKIPPED", locale))
            if csv_details and (
                csv_details["delimiterDetected"] or csv_details["encodingDetected"]
            ):
                warnings.append(_warning("CSV_DETECTION_GUESSED", locale, csv_details))
            destination = output_dir / f"{_clean_stem(item['original'])}_{index}.docx"
            _write_word(destination, parsed, locale, deadline)
            outputs.append(destination)
    except JobFailure:
        raise
    except Exception as exc:
        raise JobFailure(
            ErrorCode.OFFICE_CONVERSION_FAILED,
            "The spreadsheet could not be converted",
        ) from exc
    return OperationResult(
        _package_outputs(output_dir, outputs, "excel-to-word.zip"),
        _deduplicate_warnings(warnings),
    )
