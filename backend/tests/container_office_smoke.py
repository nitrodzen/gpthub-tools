"""Container-only Office integration smoke.

Run this script inside the built backend image. It exercises every newly
supported input family, reopens outputs with the editing libraries, and asks
LibreOffice to render each result to PDF as an independent compatibility check.
"""

from __future__ import annotations

import os
import subprocess
import tempfile
from pathlib import Path

import pikepdf
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from app.office_operations import excel_to_word, word_to_excel
from app.security import validate_signature

EXPORT_FILTERS = {
    ".doc": "doc:MS Word 97",
    ".odt": "odt:writer8",
    ".rtf": "rtf:Rich Text Format",
    ".xls": "xls:MS Excel 97",
    ".ods": "ods:calc8",
    ".pdf": "pdf",
}


def item(path: Path) -> dict[str, str]:
    return {
        "path": str(path),
        "original": path.name,
        "content_type": "application/octet-stream",
    }


def libreoffice_export(source: Path, output_dir: Path, suffix: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = output_dir / ".lo-profile"
    profile.mkdir()
    environment = {**os.environ, "HOME": str(output_dir / ".home")}
    command = [
        "libreoffice",
        "--headless",
        "--nologo",
        "--nodefault",
        "--nolockcheck",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to",
        EXPORT_FILTERS[suffix],
        "--outdir",
        str(output_dir),
        str(source),
    ]
    completed = subprocess.run(
        command,
        check=False,
        capture_output=True,
        text=True,
        timeout=90,
        env=environment,
    )
    result = output_dir / f"{source.stem}{suffix}"
    if completed.returncode or not result.is_file() or not result.stat().st_size:
        raise RuntimeError(
            f"LibreOffice export to {suffix} failed: "
            f"{completed.stdout.strip()} {completed.stderr.strip()}"
        )
    return result


def assert_libreoffice_opens(source: Path, output_dir: Path) -> None:
    rendered = libreoffice_export(source, output_dir, ".pdf")
    with pikepdf.open(rendered) as document:
        if not document.pages:
            raise AssertionError(f"LibreOffice rendered no pages for {source.suffix}")


def make_word_fixture(root: Path, picture: Path) -> Path:
    path = root / "word-source.docx"
    document = Document()
    document.add_heading("Smoke heading", level=1)
    document.add_paragraph("Editable paragraph with Unicode: Привет, мир.")
    document.add_paragraph("First list item", style="List Bullet")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(0, 2).text = "Notes"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "=2+2"
    table.cell(1, 2).text = "Wrapped\ntext"
    table.cell(2, 0).merge(table.cell(2, 1)).text = "Merged"
    document.add_picture(str(picture))
    document.save(path)
    return path


def make_excel_fixture(root: Path, picture: Path) -> Path:
    path = root / "excel-source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Amount", "Link", "Review"])
    sheet.append(["Alice", 10, "GPTHub", "Ready"])
    sheet["B3"] = "=SUM(B2,5)"
    sheet["C2"].hyperlink = "https://gpthub.ru"
    sheet["D2"].comment = Comment("Editable comment", "GPTHub")
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Merged"
    sheet.add_image(ExcelImage(str(picture)), "E2")
    hidden = workbook.create_sheet("Hidden data")
    hidden.sheet_state = "hidden"
    hidden.append(["Secret", "Included"])
    workbook.save(path)
    workbook.close()
    return path


def smoke_word_inputs(root: Path, base_docx: Path) -> None:
    sources = [base_docx]
    for suffix in (".doc", ".odt", ".rtf"):
        sources.append(libreoffice_export(base_docx, root / f"fixture-{suffix[1:]}", suffix))

    for source in sources:
        print(f"Checking Word input: {source.suffix}", flush=True)
        validate_signature(source, source.suffix.casefold(), "application/octet-stream")
        output_dir = root / f"word-output-{source.suffix[1:]}"
        output_dir.mkdir()
        result = word_to_excel([item(source)], output_dir, {"locale": "en"})
        workbook = load_workbook(result.path, data_only=False)
        try:
            if workbook.sheetnames != ["Document"]:
                raise AssertionError(
                    f"Unexpected sheets for {source.suffix}: {workbook.sheetnames}"
                )
            values = [
                cell.value
                for row in workbook.active.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            if not any("Smoke heading" in str(value) for value in values):
                raise AssertionError(f"Heading was lost for {source.suffix}")
            if not any("Alice" in str(value) for value in values):
                raise AssertionError(f"Table content was lost for {source.suffix}")
        finally:
            workbook.close()
        assert_libreoffice_opens(result.path, root / f"word-render-{source.suffix[1:]}")


def smoke_excel_inputs(root: Path, base_xlsx: Path) -> None:
    sources = [base_xlsx]
    for suffix in (".xls", ".ods"):
        sources.append(libreoffice_export(base_xlsx, root / f"fixture-{suffix[1:]}", suffix))
    csv_path = root / "spreadsheet-source.csv"
    csv_path.write_text("Name,Amount,Note\nAlice,10,Привет\nBob,20,Ready\n", encoding="utf-8")
    sources.append(csv_path)

    for source in sources:
        print(f"Checking spreadsheet input: {source.suffix}", flush=True)
        validate_signature(source, source.suffix.casefold(), "application/octet-stream")
        output_dir = root / f"excel-output-{source.suffix[1:]}"
        output_dir.mkdir()
        options = {"locale": "en", "csvDelimiter": "auto", "csvEncoding": "auto"}
        result = excel_to_word([item(source)], output_dir, options)
        document = Document(result.path)
        if not document.tables:
            raise AssertionError(f"No editable table was produced for {source.suffix}")
        table_text = "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        if "Alice" not in table_text:
            raise AssertionError(f"Spreadsheet content was lost for {source.suffix}")
        assert_libreoffice_opens(result.path, root / f"excel-render-{source.suffix[1:]}")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="gpthub-office-smoke-") as temporary:
        root = Path(temporary)
        picture = root / "picture.png"
        Image.new("RGB", (48, 32), "#39b979").save(picture, "PNG")
        base_docx = make_word_fixture(root, picture)
        base_xlsx = make_excel_fixture(root, picture)
        smoke_word_inputs(root, base_docx)
        smoke_excel_inputs(root, base_xlsx)
    print("Office container smoke passed: DOC/DOCX/ODT/RTF and XLS/XLSX/ODS/CSV")


if __name__ == "__main__":
    main()
