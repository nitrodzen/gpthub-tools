from __future__ import annotations

import asyncio
import io
import json
import shutil
import struct
import time
import zipfile
from copy import copy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from arq import Retry
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from PIL import Image

from app import main, office_operations, operations, security, worker
from app.models import ErrorCode, JobFailure, JobStatus, JobWarning, Operation
from app.office_operations import OperationResult, excel_to_word, word_to_excel
from app.security import allowed_extensions, validate_content_type, validate_signature
from app.storage import record_to_view


def _item(path: Path) -> dict:
    return {
        "path": str(path),
        "original": path.name,
        "content_type": "application/octet-stream",
    }


def _png() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (32, 20), "blue").save(output, "PNG")
    return output.getvalue()


def test_word_to_excel_keeps_one_ordered_editable_sheet_and_raster(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("=2+2")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(0, 2).text = "C"
    table.cell(1, 0).text = "parent"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 2).text = "vertical"
    table.cell(1, 2).merge(table.cell(2, 2))
    nested = table.cell(1, 0).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "nested-a"
    nested.cell(0, 1).text = "nested-b"
    image_path = tmp_path / "picture.png"
    image_path.write_bytes(_png())
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(image_path))
    document.save(source)

    output_dir = tmp_path / "output"
    output_dir.mkdir()
    result = word_to_excel([_item(source)], output_dir, {"locale": "en"})

    workbook = load_workbook(result.path, data_only=False)
    try:
        assert workbook.sheetnames == ["Document"]
        sheet = workbook.active
        assert sheet["A1"].value == "=2+2"
        assert sheet["A1"].data_type == "s"
        merges = {str(value) for value in sheet.merged_cells.ranges}
        assert {"A1:C1", "A3:B3", "C4:C5"} <= merges
        assert sheet["A2"].value is None
        values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
        nested_label = "Nested table (R2C1.1)"
        assert values.index("=2+2") < values.index("A\nB") < values.index(nested_label)
        assert "Nested table: R2C1.1" in sheet["A4"].value
        assert sheet["A4"].hyperlink is not None
        assert sheet["A4"].hyperlink.location == "'Document'!A7"
        assert "nested-a" in values
        assert len(sheet._images) == 1
    finally:
        workbook.close()
    assert [warning.code for warning in result.warnings] == ["LAYOUT_SIMPLIFIED"]


@pytest.mark.parametrize(
    ("word_alignment", "excel_alignment"),
    [
        ("start", "left"),
        ("end", "right"),
        ("both", "justify"),
        ("distribute", "justify"),
        ("center", "center"),
    ],
)
def test_word_to_excel_accepts_libreoffice_alignment_values(
    tmp_path: Path,
    word_alignment: str,
    excel_alignment: str,
) -> None:
    source = tmp_path / f"alignment-{word_alignment}.docx"
    document = Document()
    paragraph = document.add_paragraph("Aligned text")
    properties = paragraph._p.get_or_add_pPr()
    alignment = OxmlElement("w:jc")
    alignment.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", word_alignment
    )
    properties.append(alignment)
    document.save(source)

    output_dir = tmp_path / "output"
    output_dir.mkdir()
    result = word_to_excel([_item(source)], output_dir, {"locale": "en"})

    workbook = load_workbook(result.path)
    try:
        assert workbook.active["A1"].alignment.horizontal == excel_alignment
    finally:
        workbook.close()


def test_excel_to_word_marks_hidden_skips_empty_and_bands_wide_sheet(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    visible = workbook.active
    visible.title = "Wide"
    for column in range(1, 11):
        cell = visible.cell(1, column, f"Header {column}")
        font = copy(cell.font)
        font.bold = True
        cell.font = font
        visible.cell(2, column, column)
    visible["B2"] = "=1+1"
    visible["C2"].hyperlink = "https://example.com"
    visible["D2"].comment = Comment("review", "qa")
    picture = tmp_path / "picture.png"
    picture.write_bytes(_png())
    visible.add_image(ExcelImage(str(picture)), "E2")
    hidden = workbook.create_sheet("Private")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"
    workbook.create_sheet("Empty")
    workbook.save(source)
    workbook.close()

    output_dir = tmp_path / "output"
    output_dir.mkdir()
    result = excel_to_word([_item(source)], output_dir, {"locale": "en"})

    converted = Document(result.path)
    text = "\n".join(paragraph.text for paragraph in converted.paragraphs)
    assert "Wide" in text
    assert "Columns A–H" in text
    assert "Private [hidden]" in text
    assert "Empty" not in text
    assert converted.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert len(converted.inline_shapes) >= 1
    assert any(comment.text == "review" for comment in converted.comments)
    assert "FORMULAS_AS_VALUES" in {warning.code for warning in result.warnings}
    assert "OBJECTS_SKIPPED" not in {warning.code for warning in result.warnings}


def test_excel_sheets_become_individually_oriented_word_sections(tmp_path: Path) -> None:
    source = tmp_path / "sections.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Narrow 1"
    first["A1"] = "one"
    wide = workbook.create_sheet("Wide")
    for column in range(1, 9):
        wide.cell(1, column, column)
    last = workbook.create_sheet("Narrow 2")
    last["A1"] = "three"
    workbook.save(source)
    workbook.close()
    output = tmp_path / "output"
    output.mkdir()

    result = excel_to_word([_item(source)], output, {"locale": "en"})

    document = Document(result.path)
    assert len(document.sections) == 3
    assert [section.orientation for section in document.sections] == [
        WD_ORIENT.PORTRAIT,
        WD_ORIENT.LANDSCAPE,
        WD_ORIENT.PORTRAIT,
    ]


def test_csv_manual_and_auto_are_strict(tmp_path: Path) -> None:
    source = tmp_path / "people.csv"
    source.write_bytes("Имя;Город\nИван;Москва\n".encode("cp1251"))
    output = tmp_path / "output"
    output.mkdir()
    manual = excel_to_word(
        [_item(source)],
        output,
        {
            "locale": "ru",
            "csvDelimiter": "semicolon",
            "csvEncoding": "windows-1251",
        },
    )
    document = Document(manual.path)
    assert document.tables[0].cell(0, 0).text == "Имя"
    assert document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold is True

    inconsistent = tmp_path / "bad.csv"
    inconsistent.write_text("a,b\n1,2,3\n", encoding="utf-8")
    with pytest.raises(JobFailure) as failure:
        excel_to_word(
            [_item(inconsistent)],
            output,
            {"locale": "en", "csvDelimiter": "comma", "csvEncoding": "utf-8"},
        )
    assert failure.value.code == ErrorCode.CSV_PARSE_FAILED

    ambiguous = tmp_path / "ambiguous.csv"
    ambiguous.write_text("one column only", encoding="utf-8")
    with pytest.raises(JobFailure) as failure:
        excel_to_word(
            [_item(ambiguous)],
            output,
            {"locale": "en", "csvDelimiter": "auto", "csvEncoding": "auto"},
        )
    assert failure.value.code == ErrorCode.CSV_PARSE_FAILED


def test_csv_all_text_first_row_is_header_even_when_sniffer_disagrees(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "people.csv"
    source.write_text("Name,City\nIvan,Moscow\nPetr,Tula\n", encoding="utf-8")
    output = tmp_path / "output"
    output.mkdir()
    monkeypatch.setattr(office_operations.csv.Sniffer, "has_header", lambda *_args: False)

    result = excel_to_word(
        [_item(source)],
        output,
        {"locale": "en", "csvDelimiter": "comma", "csvEncoding": "utf-8"},
    )

    document = Document(result.path)
    assert document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold is True


def test_csv_rejects_sparse_rows_before_materializing_cells(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "sparse.csv"
    source.write_text("value,,,,\n", encoding="utf-8")
    output = tmp_path / "output"
    output.mkdir()
    monkeypatch.setattr(office_operations, "MAX_OUTPUT_TABLE_CELLS", 4)

    with pytest.raises(JobFailure) as failure:
        excel_to_word(
            [_item(source)],
            output,
            {"locale": "en", "csvDelimiter": "comma", "csvEncoding": "utf-8"},
        )

    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX
    assert failure.value.details["kind"] == "csvRowCells"


def test_excel_validation_is_reported_as_skipped_object(tmp_path: Path) -> None:
    source = tmp_path / "validation.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Value"
    validation = DataValidation(type="whole", operator="between", formula1="1", formula2="10")
    sheet.add_data_validation(validation)
    validation.add(sheet["A2"])
    workbook.save(source)
    workbook.close()
    output = tmp_path / "output"
    output.mkdir()

    result = excel_to_word([_item(source)], output, {"locale": "en"})

    assert "OBJECTS_SKIPPED" in {warning.code for warning in result.warnings}


def test_word_list_styles_receive_editable_prefixes(tmp_path: Path) -> None:
    source = tmp_path / "lists.docx"
    document = Document()
    document.add_paragraph("Bullet", style="List Bullet")
    document.add_paragraph("Number", style="List Number")
    document.save(source)
    output = tmp_path / "output"
    output.mkdir()

    result = word_to_excel([_item(source)], output, {"locale": "en"})

    workbook = load_workbook(result.path)
    try:
        assert workbook.active["A1"].value == "• Bullet"
        assert workbook.active["A3"].value == "1. Number"
    finally:
        workbook.close()


def test_word_header_content_emits_objects_skipped_warning(tmp_path: Path) -> None:
    source = tmp_path / "header.docx"
    document = Document()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Header content"
    document.save(source)
    output = tmp_path / "output"
    output.mkdir()

    result = word_to_excel([_item(source)], output, {"locale": "en"})

    assert "OBJECTS_SKIPPED" in {warning.code for warning in result.warnings}


def test_word_text_box_content_emits_objects_skipped_warning(tmp_path: Path) -> None:
    source = tmp_path / "textbox.docx"
    document = Document()
    paragraph = document.add_paragraph("Body")
    text_box = OxmlElement("w:txbxContent")
    box_paragraph = OxmlElement("w:p")
    box_run = OxmlElement("w:r")
    box_text = OxmlElement("w:t")
    box_text.text = "Text box content"
    box_run.append(box_text)
    box_paragraph.append(box_run)
    text_box.append(box_paragraph)
    paragraph._p.append(text_box)
    document.save(source)
    output = tmp_path / "output"
    output.mkdir()

    result = word_to_excel([_item(source)], output, {"locale": "en"})

    assert "OBJECTS_SKIPPED" in {warning.code for warning in result.warnings}


@pytest.mark.parametrize(
    "limit_name",
    ["MAX_WORD_TABLES", "MAX_WORD_TABLE_CELLS", "MAX_EXCEL_COLUMNS"],
)
def test_word_complexity_preflight_runs_before_document_dom(
    tmp_path: Path, monkeypatch, limit_name: str
) -> None:
    source = tmp_path / "table.docx"
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "value"
    document.save(source)
    monkeypatch.setattr(office_operations, limit_name, 0)

    def fail_loader(*_args, **_kwargs):
        raise AssertionError("python-docx DOM loader must not run")

    monkeypatch.setattr(office_operations, "Document", fail_loader)
    with pytest.raises(JobFailure) as failure:
        office_operations._parse_word(source, "en")
    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX


def test_word_to_excel_rejects_overlong_cell_text(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "long-paragraph.docx"
    document = Document()
    document.add_paragraph("12345")
    document.save(source)
    output = tmp_path / "output"
    output.mkdir()
    monkeypatch.setattr(office_operations, "MAX_EXCEL_CELL_CHARACTERS", 4)

    with pytest.raises(JobFailure) as failure:
        word_to_excel([_item(source)], output, {"locale": "en"})

    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX
    assert failure.value.details["kind"] == "excelCellCharacters"


@pytest.mark.parametrize("limit_name", ["MAX_SPREADSHEET_SHEETS", "MAX_SPREADSHEET_NONEMPTY_CELLS"])
def test_excel_complexity_preflight_runs_before_workbook_dom(
    tmp_path: Path, monkeypatch, limit_name: str
) -> None:
    source = tmp_path / "cell.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "value"
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(office_operations, limit_name, 0)

    def fail_loader(*_args, **_kwargs):
        raise AssertionError("openpyxl DOM loader must not run")

    monkeypatch.setattr(office_operations, "load_workbook", fail_loader)
    with pytest.raises(JobFailure) as failure:
        office_operations._parse_excel(source)
    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX


def test_excel_merge_complexity_preflight_runs_before_workbook_dom(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "large-merge.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "value"
    workbook.active.merge_cells("A1:C3")
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(office_operations, "MAX_OUTPUT_TABLE_CELLS", 4)

    def fail_loader(*_args, **_kwargs):
        raise AssertionError("openpyxl DOM loader must not run")

    monkeypatch.setattr(office_operations, "load_workbook", fail_loader)
    with pytest.raises(JobFailure) as failure:
        office_operations._parse_excel(source)

    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX
    assert failure.value.details["kind"] == "spreadsheetMergedCells"


def test_image_pixel_budgets_are_enforced_incrementally(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(office_operations, "MAX_EMBEDDED_IMAGE_PIXELS", 100)
    with pytest.raises(JobFailure) as failure:
        office_operations._normalize_raster(_png())
    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX
    assert failure.value.details["kind"] == "embeddedImagePixels"

    monkeypatch.setattr(office_operations, "MAX_EMBEDDED_IMAGE_PIXELS", 1_000)
    monkeypatch.setattr(office_operations, "MAX_TOTAL_IMAGE_PIXELS", 1_000)
    source = tmp_path / "two-images.docx"
    picture = tmp_path / "budget.png"
    picture.write_bytes(_png())
    document = Document()
    document.add_picture(str(picture))
    document.add_picture(str(picture))
    document.save(source)
    output = tmp_path / "output"
    output.mkdir()
    with pytest.raises(JobFailure) as failure:
        word_to_excel([_item(source)], output, {"locale": "en"})
    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX
    assert failure.value.details["kind"] == "imagePixels"


def test_generic_xml_element_cap_rejects_empty_structure(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "empty-elements.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("custom.xml", "<root><empty/><empty/></root>")
    monkeypatch.setattr(office_operations, "MAX_OFFICE_XML_ELEMENTS", 2)

    with pytest.raises(JobFailure) as failure:
        office_operations._safe_xml_package(source)

    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX
    assert failure.value.details["kind"] == "officeXmlElements"


@pytest.mark.parametrize(
    ("limit_name", "xml", "kind"),
    [
        ("MAX_OFFICE_XML_DEPTH", "<a><b><c/></b></a>", "officeXmlDepth"),
        (
            "MAX_OFFICE_XML_ATTRIBUTE_CHARS",
            '<a attribute="oversized"/>',
            "officeXmlAttributeCharacters",
        ),
    ],
)
def test_xml_depth_and_attribute_limits(
    tmp_path: Path, monkeypatch, limit_name: str, xml: str, kind: str
) -> None:
    source = tmp_path / "bounded.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("custom.xml", xml)
    monkeypatch.setattr(office_operations, limit_name, 2)

    with pytest.raises(JobFailure) as failure:
        office_operations._safe_xml_package(source)

    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX
    assert failure.value.details["kind"] == kind


def test_excel_cached_values_use_basic_display_number_formats() -> None:
    formula_book = Workbook()
    value_book = Workbook()
    source = formula_book.active["A1"]
    cached = value_book.active["A1"]
    source.value = "=1/2"
    source.number_format = "0%"
    cached.value = 0.5

    value, formula, _unsafe = office_operations._excel_cell_data(source, cached)

    assert formula is True
    assert value.value == "50%"
    formula_book.close()
    value_book.close()


def test_array_formula_without_cache_uses_formula_text() -> None:
    formula_book = Workbook()
    value_book = Workbook()
    source = formula_book.active["A1"]
    cached = value_book.active["A1"]
    source.value = ArrayFormula(ref="A1:A2", text="=SEQUENCE(2)")

    value, formula, unsupported = office_operations._excel_cell_data(source, cached)

    assert formula is True
    assert unsupported is False
    assert value.value == "=SEQUENCE(2)"
    formula_book.close()
    value_book.close()


def test_excel_package_shapes_are_reported_as_skipped(tmp_path: Path) -> None:
    source = tmp_path / "shape.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "value"
    workbook.save(source)
    workbook.close()
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr(
            "xl/diagrams/data1.xml",
            '<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"/>',
        )

    parsed = office_operations._parse_excel(source)

    assert parsed.unsupported_objects is True


@pytest.mark.parametrize(
    ("extension", "mime"),
    [
        (".odt", "application/vnd.oasis.opendocument.text"),
        (".ods", "application/vnd.oasis.opendocument.spreadsheet"),
    ],
)
def test_odf_script_elements_are_rejected(tmp_path: Path, extension: str, mime: str) -> None:
    source = tmp_path / f"macro{extension}"
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:'
        'xmlns:office:1.0" xmlns:script="urn:oasis:names:tc:opendocument:'
        'xmlns:script:1.0"><office:scripts><script:script '
        'script:language="ooo:Basic"/></office:scripts></office:document-content>'
    )
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", mime)
        archive.writestr("content.xml", content)

    with pytest.raises(JobFailure) as failure:
        validate_signature(source, extension)

    assert failure.value.code == ErrorCode.OFFICE_MACROS_NOT_ALLOWED


@pytest.mark.parametrize(
    ("extension", "mime"),
    [
        (".odt", "application/vnd.oasis.opendocument.text"),
        (".ods", "application/vnd.oasis.opendocument.spreadsheet"),
    ],
)
def test_odf_empty_active_content_containers_are_allowed(
    tmp_path: Path,
    extension: str,
    mime: str,
) -> None:
    source = tmp_path / f"empty-scripts{extension}"
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:'
        'xmlns:office:1.0"><office:scripts/>\n<office:event-listeners>\n'
        "</office:event-listeners></office:document-content>"
    )
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", mime)
        archive.writestr("content.xml", content)

    validate_signature(source, extension)


def test_odf_event_listener_is_rejected(tmp_path: Path) -> None:
    source = tmp_path / "event-listener.odt"
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:'
        'xmlns:office:1.0" xmlns:script="urn:oasis:names:tc:opendocument:'
        'xmlns:script:1.0"><office:event-listeners><script:event-listener '
        'script:event-name="dom:load" script:macro-name="Run"/>'
        "</office:event-listeners></office:document-content>"
    )
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)

    with pytest.raises(JobFailure) as failure:
        validate_signature(source, ".odt")

    assert failure.value.code == ErrorCode.OFFICE_MACROS_NOT_ALLOWED


@pytest.mark.parametrize(
    ("limit_name", "limit", "content"),
    [
        ("MAX_ODF_XML_MEMBER_BYTES", 32, "<root>" + "x" * 100 + "</root>"),
        ("MAX_ODF_XML_TEXT_CHARS", 8, "<root>123456789</root>"),
        ("MAX_ODF_XML_ATTRIBUTE_CHARS", 5, '<root attribute="12345"/>'),
        ("MAX_ODF_XML_DEPTH", 2, "<root><level><too-deep/></level></root>"),
    ],
)
def test_odf_xml_resource_limits_fail_before_conversion(
    tmp_path: Path,
    monkeypatch,
    limit_name: str,
    limit: int,
    content: str,
) -> None:
    source = tmp_path / "bounded.odt"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
    monkeypatch.setattr(security, limit_name, limit)

    with pytest.raises(JobFailure) as failure:
        validate_signature(source, ".odt")

    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX


def test_odf_aggregate_xml_size_is_bounded(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "aggregate.odt"
    content = "<root>content</root>"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
        archive.writestr("styles.xml", content)
    monkeypatch.setattr(security, "MAX_ODF_XML_TOTAL_BYTES", len(content) * 2 - 1)

    with pytest.raises(JobFailure) as failure:
        validate_signature(source, ".odt")

    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX


def test_libreoffice_profile_disables_links_macros_and_recalculation(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "source.ods"
    source.write_bytes(b"fixture")
    output = tmp_path / "output"
    output.mkdir()
    captured: dict[str, str | list[str]] = {}

    def fake_run(command, **_kwargs):
        registry = next(output.glob(".office-convert-*/profile/user/registrymodifications.xcu"))
        captured["registry"] = registry.read_text(encoding="utf-8")
        captured["command"] = command
        converted_dir = Path(command[command.index("--outdir") + 1])
        (converted_dir / "source.xlsx").write_bytes(b"converted")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(office_operations.subprocess, "run", fake_run)
    generated, working = office_operations._libreoffice_convert(source, output, ".xlsx")
    try:
        assert generated.is_file()
        registry = str(captured["registry"])
        assert "MacroSecurityLevel" in registry
        assert registry.count('<prop oor:name="Link"') == 2
        assert "OOXMLRecalcMode" in registry
        assert "ODFRecalcMode" in registry
        assert registry.count("<value>1</value>") >= 4
        command = captured["command"]
        assert isinstance(command, list)
        assert {"--invisible", "--norestore", "--nofirststartwizard"} <= set(command)
    finally:
        shutil.rmtree(working, ignore_errors=True)


@pytest.mark.asyncio
async def test_isolated_office_bridge_success_and_child_failure(tmp_path: Path) -> None:
    output = tmp_path / "output"
    output.mkdir()
    source = tmp_path / "people.csv"
    source.write_text("Name,City\nIvan,Moscow\n", encoding="utf-8")
    options = {"locale": "en", "csvDelimiter": "comma", "csvEncoding": "utf-8"}

    result = await operations.run_office_isolated(
        Operation.EXCEL_TO_WORD,
        [_item(source)],
        output,
        options,
    )

    assert result.path.is_file()
    assert [warning.code for warning in result.warnings] == ["LAYOUT_SIMPLIFIED"]

    invalid = tmp_path / "invalid.csv"
    invalid.write_text("a,b\n1,2,3\n", encoding="utf-8")
    with pytest.raises(JobFailure) as failure:
        await operations.run_office_isolated(
            Operation.EXCEL_TO_WORD,
            [_item(invalid)],
            output,
            options,
        )
    assert failure.value.code == ErrorCode.CSV_PARSE_FAILED


class FakeOfficeProcess:
    def __init__(self, wait_result: int = 0, wait_error: BaseException | None = None) -> None:
        self.returncode: int | None = None
        self.wait_result = wait_result
        self.wait_error = wait_error

    async def wait(self) -> int:
        if self.wait_error is not None:
            raise self.wait_error
        self.returncode = self.wait_result
        return self.wait_result


@pytest.mark.asyncio
@pytest.mark.parametrize("payload", [None, "{not-json"])
async def test_isolated_office_bridge_rejects_missing_or_malformed_payload(
    tmp_path: Path, monkeypatch, payload: str | None
) -> None:
    output = tmp_path / "output"
    output.mkdir()

    async def fake_subprocess(*args, **_kwargs):
        if payload is not None:
            Path(args[4]).write_text(payload, encoding="utf-8")
        return FakeOfficeProcess()

    monkeypatch.setattr(operations.asyncio, "create_subprocess_exec", fake_subprocess)
    with pytest.raises(JobFailure) as failure:
        await operations.run_office_isolated(
            Operation.EXCEL_TO_WORD,
            [],
            output,
            {"locale": "en"},
        )
    assert failure.value.code == ErrorCode.OFFICE_CONVERSION_FAILED


@pytest.mark.asyncio
async def test_isolated_office_bridge_rejects_result_path_escape(
    tmp_path: Path, monkeypatch
) -> None:
    output = tmp_path / "output"
    output.mkdir()
    escaped = tmp_path / "escaped.docx"
    escaped.write_bytes(b"outside")

    async def fake_subprocess(*args, **_kwargs):
        Path(args[4]).write_text(
            json.dumps({"ok": True, "path": str(escaped), "warnings": []}),
            encoding="utf-8",
        )
        return FakeOfficeProcess()

    monkeypatch.setattr(operations.asyncio, "create_subprocess_exec", fake_subprocess)
    with pytest.raises(JobFailure) as failure:
        await operations.run_office_isolated(
            Operation.EXCEL_TO_WORD,
            [],
            output,
            {"locale": "en"},
        )
    assert failure.value.code == ErrorCode.OFFICE_CONVERSION_FAILED


@pytest.mark.asyncio
@pytest.mark.parametrize("cancelled", [False, True])
async def test_isolated_office_bridge_kills_child_on_timeout_or_cancellation(
    tmp_path: Path, monkeypatch, cancelled: bool
) -> None:
    output = tmp_path / "output"
    output.mkdir()
    process = FakeOfficeProcess(
        wait_error=asyncio.CancelledError() if cancelled else None,
    )
    if not cancelled:

        async def wait_forever() -> int:
            await asyncio.sleep(60)
            return 0

        process.wait = wait_forever  # type: ignore[method-assign]

    async def fake_subprocess(*_args, **_kwargs):
        return process

    async def fake_terminate(target) -> None:
        target.returncode = -9

    terminate = AsyncMock(side_effect=fake_terminate)
    monkeypatch.setattr(operations.asyncio, "create_subprocess_exec", fake_subprocess)
    monkeypatch.setattr(operations, "terminate_office_process", terminate)
    monkeypatch.setattr(operations, "settings", SimpleNamespace(job_timeout_seconds=0.001))

    expected = asyncio.CancelledError if cancelled else JobFailure
    with pytest.raises(expected) as failure:
        await operations.run_office_isolated(
            Operation.EXCEL_TO_WORD,
            [],
            output,
            {"locale": "en"},
        )
    if not cancelled:
        assert failure.value.code == ErrorCode.TIMEOUT
    terminate.assert_awaited_once_with(process)


def test_office_extensions_mime_and_persisted_warnings() -> None:
    assert allowed_extensions(Operation.WORD_TO_EXCEL) == {".doc", ".docx", ".odt", ".rtf"}
    assert allowed_extensions(Operation.EXCEL_TO_WORD) == {".xls", ".xlsx", ".ods", ".csv"}
    validate_content_type(".xlsx", "application/octet-stream")
    validate_content_type(
        ".xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    with pytest.raises(JobFailure) as failure:
        validate_content_type(".xlsx", "image/png")
    assert failure.value.code == ErrorCode.INVALID_FILE

    record = {
        "job_id": "job",
        "operation": Operation.EXCEL_TO_WORD.value,
        "status": "succeeded",
        "created_at": "2026-01-01T00:00:00Z",
        "expires_at": "2026-01-01T01:00:00Z",
        "warnings": json.dumps(
            [
                {
                    "code": "LAYOUT_SIMPLIFIED",
                    "message": "Layout simplified",
                    "details": {"sheet": "Data"},
                }
            ]
        ),
    }
    view = record_to_view(record)
    assert view.warnings[0].code == "LAYOUT_SIMPLIFIED"
    assert view.warnings[0].details == {"sheet": "Data"}


def test_office_package_security_rejects_macros_duplicates_and_media_limit(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "safe.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "safe"
    workbook.save(source)
    workbook.close()
    validate_signature(
        source,
        ".xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    macro = tmp_path / "macro.xlsx"
    macro.write_bytes(source.read_bytes())
    with zipfile.ZipFile(macro, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(JobFailure) as failure:
        validate_signature(macro, ".xlsx")
    assert failure.value.code == ErrorCode.OFFICE_MACROS_NOT_ALLOWED

    duplicate = tmp_path / "duplicate.xlsx"
    duplicate.write_bytes(source.read_bytes())
    with pytest.warns(UserWarning, match="Duplicate name"):
        with zipfile.ZipFile(duplicate, "a") as archive:
            archive.writestr("xl/workbook.xml", b"<duplicate/>")
    with pytest.raises(JobFailure) as failure:
        validate_signature(duplicate, ".xlsx")
    assert failure.value.code == ErrorCode.INVALID_FILE

    document_path = tmp_path / "image.docx"
    picture = tmp_path / "security-picture.png"
    picture.write_bytes(_png())
    document = Document()
    document.add_picture(str(picture))
    document.save(document_path)
    monkeypatch.setattr(security, "MAX_OFFICE_MEDIA_FILES", 0)
    with pytest.raises(JobFailure) as failure:
        validate_signature(document_path, ".docx")
    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX


def test_office_macro_inspection_fails_closed(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "broken.doc"
    source.write_bytes(b"not-an-ole")

    class BrokenParser:
        def __init__(self, _path: str) -> None:
            pass

        def detect_vba_macros(self) -> bool:
            raise OSError("cannot inspect")

        def close(self) -> None:
            pass

    monkeypatch.setattr(security, "VBA_Parser", BrokenParser)
    with pytest.raises(JobFailure) as failure:
        security._office_has_macros(source)
    assert failure.value.code == ErrorCode.INVALID_FILE


def _zero_member_compressed_size(path: Path, member: str) -> None:
    payload = bytearray(path.read_bytes())
    encoded = member.encode()
    offset = 0
    while (offset := payload.find(b"PK\x03\x04", offset)) >= 0:
        name_length, extra_length = struct.unpack_from("<HH", payload, offset + 26)
        name = bytes(payload[offset + 30 : offset + 30 + name_length])
        if name == encoded:
            struct.pack_into("<I", payload, offset + 18, 0)
            break
        offset += 4 + name_length + extra_length
    offset = 0
    while (offset := payload.find(b"PK\x01\x02", offset)) >= 0:
        name_length = struct.unpack_from("<H", payload, offset + 28)[0]
        name = bytes(payload[offset + 46 : offset + 46 + name_length])
        if name == encoded:
            struct.pack_into("<I", payload, offset + 20, 0)
            break
        offset += 4 + name_length
    path.write_bytes(payload)


def test_office_package_rejects_nonempty_zero_compressed_entry(tmp_path: Path) -> None:
    source = tmp_path / "zero-size.docx"
    document = Document()
    document.add_paragraph("safe")
    document.save(source)
    with zipfile.ZipFile(source, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/media/payload.bin", b"not empty")
    _zero_member_compressed_size(source, "word/media/payload.bin")
    with pytest.raises(JobFailure) as failure:
        validate_signature(source, ".docx")
    assert failure.value.code == ErrorCode.OFFICE_TOO_COMPLEX


def _worker_record(job_id: str, status: str = "queued") -> dict[str, str]:
    return {
        "job_id": job_id,
        "operation": Operation.EXCEL_TO_WORD.value,
        "status": status,
        "total": "1",
        "files": "[]",
        "options": "{}",
        "ip_hash": "ip",
    }


class LeaseRedis:
    def __init__(self, record: dict[str, str], live_owners: set[str] | None = None) -> None:
        self.record = record
        self.live_owners = live_owners or set()
        self.active_score = 0.0
        self.job_expirations: dict[str, float] = {}

    async def hgetall(self, _key: str) -> dict[str, str]:
        return dict(self.record)

    async def set(self, *_args, **_kwargs) -> None:
        return None

    async def eval(self, script: str, key_count: int, *values: str):
        args = values[key_count:]
        if "previous_owner" in script:
            (
                queued,
                running,
                owner,
                token,
                _heartbeat_prefix,
                expires_at,
                _job_ttl,
                active_score,
                _job_id,
            ) = args
            status = self.record.get("status")
            if status == queued:
                self.record.update(
                    status=running,
                    progress="0",
                    run_owner=owner,
                    run_token=token,
                    expires_at=expires_at,
                )
                self.active_score = float(active_score)
                return worker.CLAIMED
            if status == running:
                previous_owner = self.record.get("run_owner")
                if previous_owner and previous_owner in self.live_owners:
                    return worker.CLAIM_BUSY
                self.record.update(
                    run_owner=owner,
                    run_token=token,
                    expires_at=expires_at,
                )
                self.active_score = float(active_score)
                return worker.CLAIMED
            return 0
        if "unpack(ARGV, 6)" in script:
            running, token, _ttl, _score, _job_id, *mapping = args
            if self.record.get("status") != running or self.record.get("run_token") != token:
                return 0
            self.record.update(dict(zip(mapping[::2], mapping[1::2], strict=True)))
            self.job_expirations[_job_id] = float(_score)
            return 1
        if "HDEL" in script:
            (token,) = args
            if self.record.get("run_token") != token:
                return ""
            status = self.record.get("status", "")
            self.record.pop("run_owner", None)
            self.record.pop("run_token", None)
            return status
        raise AssertionError("unexpected Lua script")


def _prepare_worker_test(tmp_path: Path, monkeypatch, job_id: str):
    root = tmp_path / job_id
    (root / "input").mkdir(parents=True)
    (root / "input" / "source.bin").write_bytes(b"input")
    (root / "output").mkdir()
    release = AsyncMock()
    monkeypatch.setattr(
        worker,
        "settings",
        SimpleNamespace(
            jobs_root=tmp_path,
            job_timeout_seconds=10,
            job_ttl_seconds=120,
            result_ttl_seconds=60,
        ),
    )
    monkeypatch.setattr(worker, "release_active_job", release)
    monkeypatch.setattr(
        worker,
        "delete_job_directory",
        lambda value: shutil.rmtree(tmp_path / value, ignore_errors=True),
    )
    monkeypatch.setattr(worker, "heartbeat", AsyncMock())
    monkeypatch.setattr(worker.metrics, "record_started", AsyncMock())
    monkeypatch.setattr(worker.metrics, "record_terminal", AsyncMock())
    return root, release


@pytest.mark.asyncio
async def test_worker_persists_operation_warnings(tmp_path: Path, monkeypatch) -> None:
    job_id = "office-job"
    root, release = _prepare_worker_test(tmp_path, monkeypatch, job_id)
    result_path = root / "output" / "result.docx"
    result_path.write_bytes(b"result")
    record = _worker_record(job_id)
    redis = LeaseRedis(record)
    warning = JobWarning(code="LAYOUT_SIMPLIFIED", message="simplified")

    async def fake_execute(*_args, **_kwargs) -> OperationResult:
        return OperationResult(result_path, [warning])

    monkeypatch.setattr(worker, "execute", fake_execute)
    monkeypatch.setattr(
        worker,
        "result_expiration",
        lambda: ("2026-01-01T01:00:00Z", 1_767_229_200.0),
    )
    await worker.run_operation({"redis": redis}, job_id)

    assert json.loads(record["warnings"]) == [
        {"code": "LAYOUT_SIMPLIFIED", "message": "simplified"}
    ]
    assert record["status"] == JobStatus.SUCCEEDED.value
    assert record["expires_at"] == "2026-01-01T01:00:00Z"
    assert redis.job_expirations[job_id] == 1_767_229_200.0
    assert not (root / "input" / "source.bin").exists()
    release.assert_awaited_once()


@pytest.mark.asyncio
async def test_live_worker_cancellation_is_cleaned_by_owner(tmp_path: Path, monkeypatch) -> None:
    job_id = "running-cancel"
    root, release = _prepare_worker_test(tmp_path, monkeypatch, job_id)
    result_path = root / "output" / "result.docx"
    result_path.write_bytes(b"partial")
    record = _worker_record(job_id)
    redis = LeaseRedis(record)

    async def cancelled_execute(*_args, **_kwargs) -> OperationResult:
        record["status"] = JobStatus.CANCELLED.value
        return OperationResult(result_path)

    monkeypatch.setattr(worker, "execute", cancelled_execute)
    await worker.run_operation({"redis": redis}, job_id)

    assert record["status"] == JobStatus.CANCELLED.value
    assert not root.exists()
    release.assert_awaited_once()


@pytest.mark.asyncio
async def test_live_owner_duplicate_is_retried_without_cleanup(tmp_path: Path, monkeypatch) -> None:
    job_id = "duplicate-worker"
    root, release = _prepare_worker_test(tmp_path, monkeypatch, job_id)
    record = _worker_record(job_id, JobStatus.RUNNING.value)
    record["run_owner"] = "other-worker"
    record["run_token"] = "other-token"
    redis = LeaseRedis(record, {"other-worker"})
    execute = AsyncMock()
    monkeypatch.setattr(worker, "execute", execute)

    with pytest.raises(Retry):
        await worker.run_operation({"redis": redis}, job_id)

    execute.assert_not_awaited()
    release.assert_not_awaited()
    assert (root / "input" / "source.bin").exists()
    assert record["run_token"] == "other-token"


@pytest.mark.asyncio
async def test_dead_owner_running_job_is_reclaimed(tmp_path: Path, monkeypatch) -> None:
    job_id = "restart-retry"
    root, release = _prepare_worker_test(tmp_path, monkeypatch, job_id)
    record = _worker_record(job_id, JobStatus.RUNNING.value)
    record.update(run_owner="dead-worker", run_token="dead-token")
    redis = LeaseRedis(record)
    result_path = root / "output" / "result.docx"
    result_path.write_bytes(b"result")
    monkeypatch.setattr(
        worker,
        "execute",
        AsyncMock(return_value=OperationResult(result_path)),
    )

    await worker.run_operation({"redis": redis}, job_id)

    assert record["status"] == JobStatus.SUCCEEDED.value
    assert "run_owner" not in record
    assert redis.active_score > time.time() + 300
    assert not (root / "input" / "source.bin").exists()
    release.assert_awaited_once()


@pytest.mark.asyncio
async def test_lost_claim_cannot_cleanup_new_owners_input(tmp_path: Path, monkeypatch) -> None:
    job_id = "lost-claim"
    root, release = _prepare_worker_test(tmp_path, monkeypatch, job_id)
    record = _worker_record(job_id)
    redis = LeaseRedis(record)
    result_path = root / "output" / "result.docx"
    result_path.write_bytes(b"stale-result")

    async def stolen_execute(*_args, **_kwargs) -> OperationResult:
        record.update(run_owner="new-worker", run_token="new-token")
        return OperationResult(result_path)

    monkeypatch.setattr(worker, "execute", stolen_execute)
    await worker.run_operation({"redis": redis}, job_id)

    assert (root / "input" / "source.bin").exists()
    assert record["run_token"] == "new-token"
    release.assert_not_awaited()


@pytest.mark.asyncio
async def test_cancelled_error_relinquishes_claim_but_preserves_retry_input(
    tmp_path: Path, monkeypatch
) -> None:
    job_id = "shutdown-retry"
    root, release = _prepare_worker_test(tmp_path, monkeypatch, job_id)
    record = _worker_record(job_id)
    redis = LeaseRedis(record)

    async def cancelled_execute(*_args, **_kwargs):
        raise asyncio.CancelledError

    monkeypatch.setattr(worker, "execute", cancelled_execute)
    with pytest.raises(asyncio.CancelledError):
        await worker.run_operation({"redis": redis}, job_id)

    assert record["status"] == JobStatus.RUNNING.value
    assert "run_owner" not in record
    assert "run_token" not in record
    assert (root / "input" / "source.bin").exists()
    release.assert_not_awaited()


@pytest.mark.asyncio
async def test_office_timeout_becomes_terminal_worker_failure(tmp_path: Path, monkeypatch) -> None:
    job_id = "inner-timeout"
    root, release = _prepare_worker_test(tmp_path, monkeypatch, job_id)
    record = _worker_record(job_id)
    redis = LeaseRedis(record)

    async def timed_out_execute(*_args, **_kwargs):
        raise JobFailure(ErrorCode.TIMEOUT, "Office conversion timed out")

    monkeypatch.setattr(worker, "execute", timed_out_execute)
    await worker.run_operation({"redis": redis}, job_id)

    assert record["status"] == JobStatus.FAILED.value
    assert record["error_code"] == ErrorCode.TIMEOUT.value
    assert not (root / "input" / "source.bin").exists()
    release.assert_awaited_once()


class CancelRedis:
    def __init__(self, record: dict[str, str], worker_wins: bool = False) -> None:
        self.record = record
        self.worker_wins = worker_wins
        self.evaluated = False
        self.deleted: list[str] = []
        self.zremmed: list[tuple[str, str]] = []

    async def eval(self, _script: str, key_count: int, *values: str) -> str:
        args = values[key_count:]
        if self.worker_wins and not self.evaluated:
            self.record.update(
                status=JobStatus.RUNNING.value,
                run_owner="worker",
                run_token="token",
            )
        self.evaluated = True
        previous = self.record.get("status", "")
        if previous in args[:2]:
            self.record["status"] = args[2]
        return previous

    async def delete(self, key: str) -> None:
        self.deleted.append(key)

    async def zrem(self, key: str, value: str) -> None:
        self.zremmed.append((key, value))


@pytest.mark.asyncio
async def test_delete_queued_job_wins_race_and_cleans_exactly_once(
    tmp_path: Path, monkeypatch
) -> None:
    job_id = "delete-queued"
    record = _worker_record(job_id)
    redis = CancelRedis(record)
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(redis=redis)))
    removed: list[str] = []
    release = AsyncMock()
    monkeypatch.setattr(main, "authorized_job", AsyncMock(return_value=dict(record)))
    monkeypatch.setattr(main, "delete_job_directory", removed.append)
    monkeypatch.setattr(main, "release_active_job", release)
    monkeypatch.setattr(main.metrics, "record_terminal", AsyncMock())

    await main.delete_job(request, job_id, "token" * 4)
    await main.delete_job(request, job_id, "token" * 4)

    assert record["status"] == JobStatus.CANCELLED.value
    assert removed == [job_id]
    release.assert_awaited_once_with(redis, "ip", job_id)
    main.metrics.record_terminal.assert_awaited_once()


@pytest.mark.asyncio
async def test_delete_running_job_loses_race_and_defers_cleanup(
    tmp_path: Path, monkeypatch
) -> None:
    job_id = "delete-running"
    record = _worker_record(job_id)
    redis = CancelRedis(record, worker_wins=True)
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(redis=redis)))
    removed: list[str] = []
    release = AsyncMock()
    monkeypatch.setattr(main, "authorized_job", AsyncMock(return_value=dict(record)))
    monkeypatch.setattr(main, "delete_job_directory", removed.append)
    monkeypatch.setattr(main, "release_active_job", release)
    monkeypatch.setattr(main.metrics, "record_terminal", AsyncMock())

    await main.delete_job(request, job_id, "token" * 4)

    assert record["status"] == JobStatus.CANCELLED.value
    assert removed == []
    release.assert_not_awaited()
    main.metrics.record_terminal.assert_awaited_once()
